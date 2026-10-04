"""More infographic components.

  diagram   node/arrow graph with automatic layering, dashed groups and `focus=` highlighting
            (repeat the same diagram on consecutive slides, each focusing one part)
  lanes     swim-lane timeline: items placed by year on parallel lanes
  mapping   problem → solution pairs
  stack     layered blocks top → bottom with arrows (model layers, tech stack)
  matrix    2 × 2 quadrant
  pyramid   hierarchy, narrow top → wide base
  funnel    narrowing stages, wide top → narrow bottom
  cycle     items around a loop
  progress  labelled progress bars (`- Label | 80%`)
"""
from __future__ import annotations

import math
import re

from lxml import etree
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from .style import E, mix, rgb

ARROW_SPLIT = re.compile(r"\s*(?:->|→)\s*")


def arg_value(info: str, key: str):
    m = re.search(rf"(?:^|\s){key}=(.+?)(?=\s+\w+=|$)", info or "")
    return m.group(1).strip() if m else None


class DiagramMixin:
    # ------------------------------------------------------------------ shared
    def arrow(self, s, x1, y1, x2, y2, color=None, width=1.5, elbow=False, begin=None, end=None, begin_idx=3,
              end_idx=1):
        kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
        c = s.shapes.add_connector(kind, E(x1), E(y1), E(x2), E(y2))
        st = c._element.find(qn("p:style"))
        if st is not None:
            c._element.remove(st)
        if begin is not None:
            c.begin_connect(begin, begin_idx)
        if end is not None:
            c.end_connect(end, end_idx)
        c.line.color.rgb = rgb(color or self.black)
        c.line.width = Pt(width)
        ln = c._element.spPr.find(qn("a:ln"))
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
        return c

    def node_box(self, s, x, y, w, h, title, caption="", fill=None, faded=False, focus=False, size=14):
        fill = fill or self.white
        if faded:
            sp = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE if self.radius > 0 else MSO_SHAPE.RECTANGLE, x, y, w, h,
                            fill=self.ground, line=mix(self.muted, self.ground, 0.45), lw=1.0)
            if self.radius > 0:
                sp.adjustments[0] = min(0.5, self.radius / min(w, h))
            ink, sub = mix(self.muted, self.ground, 0.25), mix(self.muted, self.ground, 0.45)
        else:
            sp = self.block(s, x, y, w, h, fill, thin=True)
            ink = self.card_ink if fill == self.white else self.on(fill)
            sub = mix(ink, fill, 0.35)
        lines = [{"text": title, "size": size, "bold": True, "color": ink}]
        if caption:
            lines.append({"text": caption, "size": max(9, size - 4), "color": sub})
        self.text(s, 0, 0, 0, 0, lines, size, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, target=sp, margin=0.06)
        return sp

    # ------------------------------------------------------------------ diagram
    def comp_diagram(self, s, d, x, y, w, h):
        info = d.get("info", "")
        nodes, captions, edges, groups = [], {}, [], []

        def add(n):
            if n and n not in nodes:
                nodes.append(n)

        for it in d["items"]:
            parts = ARROW_SPLIT.split(it["title"])
            if len(parts) > 1:
                for a, b in zip(parts, parts[1:]):
                    add(a)
                    add(b)
                    edges.append((a, b))
            else:
                add(it["title"])
                if it["desc"]:
                    captions[it["title"]] = it["desc"]
        for ln in d["lines"]:
            m = re.match(r"group\s+(.+?)\s*[:：]\s*(.+)$", ln, re.I)
            if m:
                members = [p.strip() for p in m.group(2).split(",") if p.strip()]
                groups.append((m.group(1).strip(), members))
                for p in members:
                    add(p)
            elif "->" in ln or "→" in ln:
                parts = ARROW_SPLIT.split(ln)
                for a, b in zip(parts, parts[1:]):
                    add(a)
                    add(b)
                    edges.append((a, b))
        if not nodes:
            return
        focus = {f.strip() for f in (arg_value(info, "focus") or "").split(",") if f.strip()}
        vertical = (arg_value(info, "dir") or "lr").lower() in ("tb", "td", "down")

        # longest-path layering (cycles: later edges are ignored)
        layer = {n: 0 for n in nodes}
        for _ in range(len(nodes)):
            changed = False
            for a, b in edges:
                if layer[b] < layer[a] + 1 and layer[a] + 1 < len(nodes):
                    layer[b] = layer[a] + 1
                    changed = True
            if not changed:
                break
        n_layers = max(layer.values()) + 1
        cols = [[n for n in nodes if layer[n] == k] for k in range(n_layers)]
        most = max(len(c) for c in cols)
        if vertical:
            row_h = h / n_layers
            nh = min(0.9, row_h - 0.45)
            col_w = w / most
            nw = min(2.6, col_w - 0.35)
        else:
            col_w = w / n_layers
            nw = min(2.5, col_w - 0.55)
            nh = min(1.05, (h - 0.3) / most - 0.35)
        pos = {}
        for k, col in enumerate(cols):
            m = len(col)
            for j, n in enumerate(col):
                if vertical:
                    cx = x + (w - m * (w / most)) / 2 + (j + 0.5) * (w / most)
                    cy = y + (k + 0.5) * row_h
                else:
                    cx = x + (k + 0.5) * col_w
                    span = (h - 0.3) / m
                    cy = y + 0.3 + (j + 0.5) * span
                pos[n] = (cx - nw / 2, cy - nh / 2)
        for label, members in groups:  # dashed group frames go first (behind)
            pts = [pos[mbr] for mbr in members if mbr in pos]
            if not pts:
                continue
            gx0 = min(p[0] for p in pts) - 0.22
            gy0 = min(p[1] for p in pts) - 0.38
            gx1 = max(p[0] for p in pts) + nw + 0.22
            gy1 = max(p[1] for p in pts) + nh + 0.2
            g = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, gx0, gy0, gx1 - gx0, gy1 - gy0, fill=None,
                           line=self.muted, lw=1.0)
            g.adjustments[0] = 0.05
            self.dash(g, "dash")
            self.text(s, gx0 + 0.12, gy0 + 0.04, gx1 - gx0, 0.3, label, 10, color=self.muted, font="label")
        shapes = {}
        for k, n in enumerate(nodes):
            faded = bool(focus) and n not in focus
            hot = n in focus
            fill = self.pal(0) if hot else self.white
            shapes[n] = self.node_box(s, *pos[n], nw, nh, n, captions.get(n, ""), fill=fill, faded=faded, focus=hot)
        for a, b in edges:
            dim = bool(focus) and not (a in focus or b in focus)
            col = mix(self.muted, self.ground, 0.4) if dim else self.black
            if vertical:
                self.arrow(s, 0, 0, 1, 1, color=col, width=1.25, elbow=True, begin=shapes[a], end=shapes[b],
                           begin_idx=2, end_idx=0)
            else:
                self.arrow(s, 0, 0, 1, 1, color=col, width=1.25, elbow=True, begin=shapes[a], end=shapes[b])

    # ------------------------------------------------------------------ lanes
    def comp_lanes(self, s, d, x, y, w, h):
        items = d["items"]
        if not items:
            return
        info = d.get("info", "")

        def year(t):
            m = re.search(r"\d{4}(?:\.\d+)?", t)
            return float(m.group(0)) if m else None

        lanes = []
        for it in items:
            ln = it["attrs"].get("lane", "")
            if ln not in lanes:
                lanes.append(ln)
        years = [year(it["title"]) for it in items if year(it["title"]) is not None]
        rng = re.search(r"(\d{4})\s*[-–~]\s*(\d{4})", info or "")
        y0, y1 = (float(rng.group(1)), float(rng.group(2))) if rng else (min(years), max(years))
        if y1 <= y0:
            y1 = y0 + 1
        label_w = 1.5
        axis_h = 0.45
        lane_h = (h - axis_h - 0.15 * (len(lanes) - 1)) / len(lanes)
        tx0, tx1 = x + label_w + 0.35, x + w - 0.5
        X = lambda v: tx0 + (v - y0) / (y1 - y0) * (tx1 - tx0)  # noqa: E731
        for k, lane in enumerate(lanes):
            ly = y + k * (lane_h + 0.15)
            band = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE if self.radius > 0 else MSO_SHAPE.RECTANGLE, x, ly, w,
                              lane_h, fill=mix(self.pal(k), self.ground, 0.7))
            if self.radius > 0:
                band.adjustments[0] = 0.08
            self.text(s, x + 0.15, ly, label_w, lane_h, lane, 14, bold=True,
                      color=mix(self.pal(k), "000000", 0.35) if self.luminance(self.pal(k)) > 0.5 else self.pal(k),
                      anchor=MSO_ANCHOR.MIDDLE, font="heading")
            rows: list[list[tuple[float, float]]] = []
            lane_items = [it for it in items if it["attrs"].get("lane", "") == lane]
            for it in sorted(lane_items, key=lambda i: year(i["title"]) or y0):
                name = it["desc"] or it["title"]
                pw = self.measure(name, 12) + 0.45
                cx = X(year(it["title"]) or y0)
                left = min(max(tx0 - 0.2, cx - pw / 2), tx1 + 0.3 - pw)
                r = next((i for i, row in enumerate(rows) if all(left > b + 0.08 or left + pw < a - 0.08
                                                                   for a, b in row)), len(rows))
                if r == len(rows):
                    rows.append([])
                rows[r].append((left, left + pw))
                slot = min(0.66, lane_h / max(2, len(rows) + 0.5))
                py = ly + 0.18 + r * slot
                sp = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, left, py, pw, 0.34, fill=self.white,
                                line=mix(self.muted, self.ground, 0.2), lw=0.75)
                self.text(s, 0, 0, 0, 0, name, 12, bold=True, color=self.on(self.white), align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE, target=sp)
                if it["attrs"].get("note"):
                    self.text(s, left, py + 0.34, pw, 0.22, it["attrs"]["note"], 9, color=self.muted,
                              align=PP_ALIGN.CENTER)
        ay = y + h - axis_h + 0.1
        self.line(s, tx0 - 0.2, ay, tx1 + 0.2, ay, color=self.muted, width=1)
        step = max(1, round((y1 - y0) / 10))
        v = y0
        while v <= y1 + 0.01:
            self.line(s, X(v), ay - 0.05, X(v), ay + 0.05, color=self.muted, width=1)
            self.text(s, X(v) - 0.4, ay + 0.05, 0.8, 0.25, f"{int(v)}", 9, color=self.muted, align=PP_ALIGN.CENTER)
            v += step

    # ------------------------------------------------------------------ mapping
    def comp_mapping(self, s, d, x, y, w, h):
        pairs = []
        for it in d["items"]:
            parts = ARROW_SPLIT.split(it["title"])
            if len(parts) >= 2:
                pairs.append((parts[0], parts[-1]))
            elif it["desc"]:
                pairs.append((it["title"], it["desc"]))
        if not pairs:
            return
        neutral = "neutral" in d["args"]
        lf = self.pal(0) if neutral else self.tones["negative"]
        rf = self.pal(1) if neutral else self.tones["positive"]
        n = len(pairs)
        gap = 0.3
        rh = min(0.72, (h - gap * (n - 1)) / n)
        y0 = y + (h - (rh * n + gap * (n - 1))) / 2
        cw = w * 0.4
        for k, (a, b) in enumerate(pairs):
            ry = y0 + k * (rh + gap)
            for bx, txt, fill in ((x, a, lf), (x + w - cw, b, rf)):
                sp = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, bx, ry, cw, rh, fill=fill)
                if self.surface in ("brutal", "pixel"):
                    sp.line.color.rgb = rgb(self.black)
                    sp.line.width = Pt(self.tw)
                elif self.surface in ("soft", "clay", "glass"):
                    self.soft_shadow(sp, blur=8, dist=3, alpha=0.18)
                size = 16
                while size > 11 and self.measure(txt, size) > cw - 0.4:
                    size -= 1
                self.text(s, 0, 0, 0, 0, txt, size, bold=True, color=self.on(fill), align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE, target=sp, margin=0.1)
            self.arrow(s, x + cw + 0.25, ry + rh / 2, x + w - cw - 0.25, ry + rh / 2, color=self.title_ink,
                       width=1.5)

    # ------------------------------------------------------------------ stack
    def comp_stack(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        aw = 0.32
        bh = min(0.85, (h - aw * (n - 1)) / n)
        bw = min(w, 7.5)
        has_desc = any(it["desc"] for it in items)
        bx = x + (w - bw) / 2 if not has_desc else x
        y0 = y + (h - (bh * n + aw * (n - 1))) / 2
        for k, it in enumerate(items):
            by = y0 + k * (bh + aw)
            fill = self.pal(k)
            sp = self.block(s, bx, by, bw, bh, fill, thin=True)
            self.text(s, 0, 0, 0, 0, it["title"], 17, bold=True, color=self.on(fill), align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE, target=sp)
            if it["desc"]:
                self.text(s, bx + bw + 0.35, by, w - bw - 0.35, bh, it["desc"], 14, color=self.ink,
                          anchor=MSO_ANCHOR.MIDDLE)
            if k < n - 1:
                self.shape(s, MSO_SHAPE.DOWN_ARROW, bx + bw / 2 - 0.14, by + bh + 0.04, 0.28, aw - 0.08,
                           fill=self.title_ink)

    # ------------------------------------------------------------------ matrix
    def comp_matrix(self, s, d, x, y, w, h):
        items = (d["items"] + [None] * 4)[:4]
        info = d.get("info", "")
        xl, yl = arg_value(info, "x") or "", arg_value(info, "y") or ""
        ax = 0.45
        gx, gy, gw, gh = x + ax, y, w - ax, h - ax
        g = 0.18
        cw, ch = (gw - g) / 2, (gh - g) / 2
        for k, it in enumerate(items):
            r, c = divmod(k, 2)
            fill = mix(self.pal(k), self.ground, 0.55 if k != 1 else 0.15)
            bx, by = gx + c * (cw + g), gy + r * (ch + g)
            self.block(s, bx, by, cw, ch, fill, thin=True)
            if it:
                lines = [{"text": it["title"], "size": 20, "bold": True, "color": self.on(fill), "font": "heading",
                          "space_after": 4}]
                if it["desc"]:
                    lines.append({"text": it["desc"], "size": 14, "color": self.on(fill)})
                lines += [{"text": "■ " + ch_, "size": 13, "color": self.on(fill)} for ch_ in it["children"]]
                self.text(s, bx + 0.3, by + 0.25, cw - 0.6, ch - 0.5, lines, 20)
        self.arrow(s, x + 0.15, y + h - ax + 0.1, x + 0.15, y, color=self.title_ink, width=2)
        self.arrow(s, x + 0.15, y + h - 0.12, x + w, y + h - 0.12, color=self.title_ink, width=2)
        if yl:
            t = self.text(s, x - 0.32, y, 0.4, gh, yl, 12, bold=True, color=self.title_ink, align=PP_ALIGN.CENTER,
                          font="label", anchor=MSO_ANCHOR.MIDDLE)
            t.text_frame._txBody.bodyPr.set("vert", "vert270")
        if xl:
            self.text(s, gx, y + h - 0.1, gw, 0.35, xl, 12, bold=True, color=self.title_ink, align=PP_ALIGN.CENTER,
                      font="label")

    # ------------------------------------------------------------------ pyramid / funnel
    def _levels(self, s, d, x, y, w, h, invert):
        items = d["items"]
        n = len(items)
        if not n:
            return
        has_desc = any(it["desc"] or it["children"] for it in items)
        pw = min(w * (0.55 if has_desc else 0.9), 7.0)
        g = 0.06
        lh = (h - g * (n - 1)) / n
        for k, it in enumerate(items):
            i = (n - 1 - k) if invert else k
            top_w = pw * (i / n) if not invert else pw * ((i + 1) / n)
            bot_w = pw * ((i + 1) / n) if not invert else pw * (i / n)
            top_w = max(top_w, pw * 0.18) if invert else top_w
            bot_w = max(bot_w, pw * 0.18) if invert else bot_w
            ly = y + k * (lh + g)
            big = max(top_w, bot_w)
            bx = x + (pw - big) / 2
            fill = self.pal(k)
            if k == 0 and not invert:
                sp = self.shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, bx, ly, big, lh, fill=fill)
            else:
                sp = self.shape(s, MSO_SHAPE.TRAPEZOID, bx, ly, big, lh, fill=fill)
                sp.adjustments[0] = min(1.0, abs(bot_w - top_w) / 2 / max(0.05, min(big, lh)))
                if invert:
                    sp._element.spPr.find(qn("a:xfrm")).set("flipV", "1")
            if self.surface in ("brutal", "pixel", "outline", "sketch"):
                sp.line.color.rgb = rgb(self.black)
                sp.line.width = Pt(self.tw)
            elif self.surface in ("soft", "clay", "glass", "neu"):
                self.soft_shadow(sp, blur=8, dist=3, alpha=0.2)
            ty = ly + (lh * 0.25 if (k == 0 and not invert) else 0)
            self.text(s, x, ty, pw, lh - (lh * 0.25 if (k == 0 and not invert) else 0), it["title"], 15, bold=True,
                      color=self.on(fill), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if has_desc:
                lines = [{"text": it["desc"], "size": 14, "color": self.ink}] if it["desc"] else []
                lines += [{"text": "■ " + c, "size": 12, "color": self.muted} for c in it["children"]]
                self.text(s, x + pw + 0.4, ly, w - pw - 0.4, lh, lines, 14, anchor=MSO_ANCHOR.MIDDLE)
                self.line(s, x + (pw + big) / 2 + 0.05, ly + lh / 2, x + pw + 0.3, ly + lh / 2,
                          color=mix(self.muted, self.ground, 0.3), width=0.75)

    def comp_pyramid(self, s, d, x, y, w, h):
        self._levels(s, d, x, y, w, h, invert=False)

    def comp_funnel(self, s, d, x, y, w, h):
        self._levels(s, d, x, y, w, h, invert=True)

    # ------------------------------------------------------------------ cycle
    def comp_cycle(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        center = " ".join(d["args"])
        r = min(w, h) / 2 - 0.75
        cx, cy = x + w / 2, y + h / 2
        nd = min(1.5, 2 * math.pi * r / n * 0.62)
        pts = []
        for k in range(n):
            a = -math.pi / 2 + 2 * math.pi * k / n
            pts.append((cx + r * math.cos(a), cy + r * math.sin(a)))
        for k in range(n):  # arrows along the ring, drawn first
            (x1, y1), (x2, y2) = pts[k], pts[(k + 1) % n]
            dx, dy = x2 - x1, y2 - y1
            dist = math.hypot(dx, dy)
            off = nd / 2 + 0.08
            self.arrow(s, x1 + dx / dist * off, y1 + dy / dist * off, x2 - dx / dist * off, y2 - dy / dist * off,
                       color=self.title_ink, width=2)
        for k, (px, py) in enumerate(pts):
            fill = self.pal(k)
            sp = self.block(s, px - nd / 2, py - nd / 2, nd, nd, fill, thin=True, kind=MSO_SHAPE.OVAL)
            it = items[k]
            lines = [{"text": it["title"], "size": 13, "bold": True, "color": self.on(fill)}]
            if it["desc"]:
                lines.append({"text": it["desc"], "size": 10, "color": self.on(fill)})
            self.text(s, 0, 0, 0, 0, lines, 13, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, target=sp,
                      margin=0.08)
        if center:
            self.text(s, cx - r * 0.55, cy - 0.4, r * 1.1, 0.8, center, 18, bold=True, color=self.title_ink,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font="heading")

    # ------------------------------------------------------------------ progress
    def comp_progress(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        rh = min(0.8, h / n)
        y0 = y + (h - rh * n) / 2
        lw = min(3.6, w * 0.35)
        track_w = w - lw - 1.2
        for k, it in enumerate(items):
            ry = y0 + k * rh
            m = re.search(r"([\d.]+)\s*%", it["desc"] or it["title"])
            pct = max(0.0, min(100.0, float(m.group(1)))) if m else 0.0
            self.text(s, x, ry, lw, rh, it["title"], 16, bold=True, color=self.title_ink, anchor=MSO_ANCHOR.MIDDLE)
            bh = min(0.32, rh * 0.45)
            by = ry + (rh - bh) / 2
            track = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + lw, by, track_w, bh,
                               fill=mix(self.muted, self.ground, 0.75))
            track.adjustments[0] = 0.5 if self.radius > 0 or self.surface not in ("brutal", "pixel") else 0
            if pct > 0:
                bar = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + lw, by, track_w * pct / 100, bh, fill=self.pal(k))
                bar.adjustments[0] = track.adjustments[0]
                if self.surface in ("brutal", "pixel"):
                    bar.line.color.rgb = rgb(self.black)
                    bar.line.width = Pt(self.tw)
            self.text(s, x + lw + track_w + 0.15, ry, 1.05, rh, f"{pct:g}%", 16, bold=True, color=self.title_ink,
                      anchor=MSO_ANCHOR.MIDDLE, font="display")
