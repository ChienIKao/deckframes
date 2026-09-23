"""Template engine: fills a .pptx/.potx master's placeholders (no nav bar / sub-TOC)."""
from __future__ import annotations

import io
import json
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from ..layout import EMU_PER_PT, TEXT_KINDS, fit_size, merge_bullets, plan, text_units
from ..markdown import SlideSpec, disp_width, inline_runs, plain

# --------------------------------------------------------------------------- pptx helpers

TEMPLATE_CT = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
PRES_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"


def open_presentation(path: Path | None):
    if path is None:
        return Presentation()
    data = path.read_bytes()
    if path.suffix.lower() == ".potx":
        src = zipfile.ZipFile(io.BytesIO(data))
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                content = src.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    content = content.replace(TEMPLATE_CT.encode(), PRES_CT.encode())
                dst.writestr(item, content)
        data = buf.getvalue()
    return Presentation(io.BytesIO(data))


def remove_all_slides(prs):
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)


def hex_rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


def is_dark(h: str | None) -> bool:
    if not h:
        return False
    r, g, b = (int(h.lstrip("#")[k:k + 2], 16) for k in (0, 2, 4))
    return 0.299 * r + 0.587 * g + 0.114 * b < 128


def set_font(run, family=None, ea=None, size=None, color=None, bold=None, italic=None):
    f = run.font
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color:
        f.color.rgb = hex_rgb(color)
    if family:
        f.name = family
    if ea:
        rPr = run._r.get_or_add_rPr()
        el = rPr.find(qn("a:ea"))
        if el is None:
            el = etree.SubElement(rPr, qn("a:ea"))
            latin = rPr.find(qn("a:latin"))
            anchor = None
            for tag in ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"):
                anchor = rPr.find(qn(tag))
                if anchor is not None:
                    break
            rPr.remove(el)
            if latin is not None:
                latin.addnext(el)
            elif anchor is not None:
                anchor.addprevious(el)
            else:
                rPr.append(el)
        el.set("typeface", ea)


BU_TAGS = ("a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip")
PPR_TAIL = ("a:tabLst", "a:defRPr", "a:extLst")


def set_bullet(p, mode, level=0):
    """mode: inherit | none | num | char"""
    if mode == "inherit":
        return
    pPr = p._p.get_or_add_pPr()
    for tag in BU_TAGS:
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if mode == "none":
        el = etree.Element(qn("a:buNone"))
        pPr.set("marL", "0")
        pPr.set("indent", "0")
    elif mode == "num":
        el = etree.Element(qn("a:buAutoNum"), type="arabicPeriod")
    else:
        el = etree.Element(qn("a:buChar"), char="•")
        pPr.set("marL", str(int(Inches(0.3) * (level + 1))))
        pPr.set("indent", str(-int(Inches(0.25))))
    anchor = None
    for tag in PPR_TAIL:
        anchor = pPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(el)
    else:
        pPr.append(el)


def box(shape):
    return (shape.left, shape.top, shape.width, shape.height)


# --------------------------------------------------------------------------- renderer

ROLE_KEYWORDS = {
    "cover": ["title slide", "標題投影片", "标题幻灯片", "封面", "cover"],
    "section": ["section", "章節", "章节", "節標題", "节标题", "區段", "divider"],
    "content": ["title and content", "標題及內容", "标题和内容", "標題及物件", "content"],
    "two_content": ["two content", "兩個內容", "两栏内容", "两项内容", "two column"],
    "title_only": ["title only", "只有標題", "仅标题", "僅標題"],
    "closing": ["closing", "end", "結尾", "结束", "thank"],
}
BODY_TYPES = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
META_TYPES = (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)


def ph_types(layout):
    return [ph.placeholder_format.type for ph in layout.placeholders]


def detect_layouts(prs, mapping: dict):
    layouts = list(prs.slide_layouts)
    by_name = {l.name: l for l in layouts}
    found = {}
    for role, ref in (mapping or {}).items():
        if isinstance(ref, int) or (isinstance(ref, str) and ref.isdigit()):
            found[role] = layouts[int(ref)]
        elif ref in by_name:
            found[role] = by_name[ref]
        else:
            print(f"warning: layout '{ref}' for role '{role}' not found", file=sys.stderr)

    def structural(role, l):
        t = ph_types(l)
        bodies = sum(1 for x in t if x in BODY_TYPES)
        has_title = any(x in TITLE_TYPES for x in t)
        return {
            "cover": PP_PLACEHOLDER.CENTER_TITLE in t,
            "content": has_title and bodies == 1,
            "two_content": has_title and bodies == 2,
            "title_only": has_title and bodies == 0,
            "section": has_title and bodies == 1,
        }.get(role, False)

    for role, kws in ROLE_KEYWORDS.items():
        if role in found:
            continue
        for l in layouts:
            if any(k in l.name.lower() for k in kws) and (role == "closing" or structural(role, l) or role == "section"):
                found[role] = l
                break
        else:
            for l in layouts:
                if role != "closing" and structural(role, l):
                    found[role] = l
                    break
    found.setdefault("content", layouts[min(1, len(layouts) - 1)])
    found.setdefault("cover", layouts[0])
    found.setdefault("section", found["cover"])
    found.setdefault("title_only", found["content"])
    found.setdefault("agenda", found["content"])
    found.setdefault("closing", found["section"])
    return found


def scale_to_widescreen(prs):
    """Default python-pptx template is 4:3; stretch every master/layout placeholder to 16:9."""
    ratio = Inches(13.333) / prs.slide_width
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    master = prs.slide_master
    for shp in list(master.shapes) + [s for l in prs.slide_layouts for s in l.shapes]:
        # only shapes with their own xfrm; inherited geometry follows the (already scaled) master
        for xfrm in shp._element.xpath("./p:spPr/a:xfrm"):
            off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
            off.set("x", str(int(int(off.get("x")) * ratio)))
            ext.set("cx", str(int(int(ext.get("cx")) * ratio)))


class Renderer:
    def __init__(self, prs, cfg, theme, base_dir: Path):
        self.prs, self.cfg, self.theme, self.base = prs, cfg, theme, base_dir
        self.layouts = detect_layouts(prs, cfg.get("layouts", {}))
        self.warnings = []
        t = theme or {}
        self.fonts = t.get("fonts", {})
        self.colors = t.get("colors", {})
        self.bgs = t.get("backgrounds", {})
        self.sizes = {"title": None, "cover_title": None, "body_max": 24, "body_min": 14,
                      "code_max": 20, "code_min": 10, "table_max": 20, "table_min": 10, "caption": 12}
        self.sizes.update(cfg.get("sizes", {}))
        self.sizes.update(t.get("sizes", {}))
        self.title_align = t.get("title_align") or cfg.get("title_align")

    # ---- styling
    def dark(self, role):
        return is_dark(self.bgs.get(role))

    def color(self, role, key):
        if not self.colors:
            return None
        if self.dark(role):
            return self.colors.get(f"{key}_on_dark") or self.colors.get("text_on_dark")
        return self.colors.get(key)

    def write_runs(self, p, text, role, size=None, bold=None, italic=None, color_key="text", heading=False):
        fam = self.fonts.get("heading" if heading else "body")
        ea = self.fonts.get("ea_heading" if heading else "ea_body") or self.fonts.get("ea")
        for chunk, fmt in inline_runs(text):
            if not chunk:
                continue
            r = p.add_run()
            r.text = chunk
            c = self.color(role, "accent") if fmt.get("bold") and self.colors.get("emphasis_accent") else self.color(role, color_key)
            if fmt.get("code"):
                set_font(r, family=self.fonts.get("code", "Consolas"), ea=ea, size=size, color=c)
            else:
                set_font(r, family=fam, ea=ea, size=size, color=c,
                         bold=True if fmt.get("bold") else bold, italic=True if fmt.get("italic") else italic)
            if fmt.get("link"):
                r.hyperlink.address = fmt["link"]

    def fill_paras(self, tf, paras, role, size, textbox=False):
        tf.word_wrap = True
        # drop every existing paragraph but the first, clear the first
        for extra in tf.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        first = tf.paragraphs[0]
        for r in list(first.runs):
            r._r.getparent().remove(r._r)
        for n, pd in enumerate(paras):
            p = first if n == 0 else tf.add_paragraph()
            p.level = pd.get("level", 0)
            mode = pd.get("bullet", "inherit")
            if textbox and mode == "inherit":
                mode = "char"
            set_bullet(p, mode, pd.get("level", 0))
            if pd.get("align"):
                p.alignment = pd["align"]
            s = max(10, round(size * pd.get("scale", 1.0)))
            p.space_after = Pt(round(s * 0.35))
            self.write_runs(p, pd["text"], role, size=s, bold=pd.get("bold"), italic=pd.get("italic"),
                            color_key=pd.get("color", "text"))

    def set_title(self, slide, text, role):
        ph = slide.shapes.title
        if ph is None:
            return
        tf = ph.text_frame
        tf.text = ""
        p = tf.paragraphs[0]
        size = self.sizes.get("cover_title") if role in ("cover", "closing") else self.sizes.get("title")
        if self.title_align and role in ("content", "agenda"):
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[self.title_align]
        self.write_runs(p, text, role, size=size, color_key="title", heading=True,
                        bold=(self.theme or {}).get("title_bold"))

    def background(self, slide, role):
        c = self.bgs.get(role) or (self.bgs.get("content") if role in ("agenda",) else None)
        if c:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = hex_rgb(c)

    def new_slide(self, role, layout_role=None):
        slide = self.prs.slides.add_slide(self.layouts[layout_role or role])
        self.background(slide, role)
        return slide

    def placeholders(self, slide, types):
        return sorted([ph for ph in slide.placeholders if ph.placeholder_format.type in types],
                      key=lambda p: (p.left or 0, p.top or 0))

    def cleanup(self, slide):
        for ph in list(slide.placeholders):
            if ph.placeholder_format.type in META_TYPES:
                continue
            if ph.has_text_frame and ph.text_frame.text.strip():
                continue
            ph._element.getparent().remove(ph._element)

    def notes(self, slide, notes):
        if notes:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(notes)

    def content_area(self):
        lay = self.layouts["content"]
        bodies = [ph for ph in lay.placeholders if ph.placeholder_format.type in BODY_TYPES]
        if bodies and bodies[0].width:
            return box(bodies[0])
        W, H = self.prs.slide_width, self.prs.slide_height
        return (int(W * 0.06), int(H * 0.22), int(W * 0.88), int(H * 0.70))

    def split_to_fit(self, spec: SlideSpec):
        """Halve a content slide until its text fits the real box at `split_below` pt."""
        if spec.role != "content":
            return [spec]
        text = [b for b in spec.blocks if b.kind in TEXT_KINDS]
        vis = [b for b in spec.blocks if b.kind not in TEXT_KINDS]
        units = [u for b in text for u, _ in text_units(b)]
        if len(units) < 2:
            return [spec]
        _, _, w, h = self.content_area()
        if vis:
            w = int(w * float(self.cfg.get("text_ratio", 0.45)))
        floor = self.sizes.get("split_below", 16)
        if fit_size(self.block_paras(text), w, h, floor, floor)[1]:
            return [spec]
        mid = len(units) // 2
        suffix = self.cfg.get("continued_suffix", "（續）")
        cont = spec.title if spec.title.endswith(suffix) or not spec.title else spec.title + suffix
        a = SlideSpec("content", spec.title, blocks=merge_bullets(units[:mid]) + vis, notes=spec.notes)
        b = SlideSpec("content", cont, blocks=merge_bullets(units[mid:]))
        return self.split_to_fit(a) + self.split_to_fit(b)

    # ---- slide kinds
    def render(self, spec: SlideSpec):
        getattr(self, f"render_{spec.role}")(spec)

    def render_cover(self, spec):
        s = self.new_slide("cover")
        self.set_title(s, spec.title, "cover")
        subs = self.placeholders(s, (PP_PLACEHOLDER.SUBTITLE,) + BODY_TYPES)
        if subs and spec.subtitle:
            paras = [{"text": ln, "bullet": "none", "color": "muted" if k else "text"}
                     for k, ln in enumerate(spec.subtitle.split("\n"))]
            size, _ = fit_size(paras, subs[0].width, subs[0].height, self.sizes.get("subtitle", 22), 12)
            self.fill_paras(subs[0].text_frame, paras, "cover", size)
            for p in subs[0].text_frame.paragraphs:
                p.alignment = None  # inherit layout alignment
        self.notes(s, spec.notes)
        self.cleanup(s)

    def render_section(self, spec):
        s = self.new_slide("section")
        self.set_title(s, spec.title, "section")
        bodies = self.placeholders(s, BODY_TYPES + (PP_PLACEHOLDER.SUBTITLE,))
        text = "\n".join(x for x in (spec.label, spec.subtitle) if x)
        if bodies and text:
            paras = [{"text": ln, "bullet": "none", "color": "accent" if k == 0 and spec.label else "muted"}
                     for k, ln in enumerate(text.split("\n"))]
            self.fill_paras(bodies[0].text_frame, paras, "section", self.sizes.get("section_label", 24))
        self.notes(s, spec.notes)
        self.cleanup(s)

    def render_closing(self, spec):
        s = self.new_slide("closing")
        self.set_title(s, spec.title, "closing")
        self.cleanup(s)

    def render_agenda(self, spec):
        s = self.new_slide("agenda")
        self.set_title(s, spec.title, "agenda")
        items = spec.blocks[0].data
        paras = [{"text": f"{k + 1:02d}　{t}", "bullet": "none"} for k, t in enumerate(items)]
        bodies = self.placeholders(s, BODY_TYPES)
        if bodies:
            w, h = bodies[0].width, bodies[0].height
            size, _ = fit_size(paras, w, h, self.sizes["body_max"] + 8, self.sizes["body_min"])
            self.fill_paras(bodies[0].text_frame, paras, "agenda", size)
        self.cleanup(s)

    def block_paras(self, blocks):
        paras = []
        for b in blocks:
            if b.kind == "bullets":
                for it in b.data:
                    paras.append({"text": it["text"], "level": it["level"],
                                  "bullet": "num" if it["ordered"] else "inherit",
                                  "scale": max(0.7, 1 - 0.1 * it["level"])})
            elif b.kind == "para":
                paras.append({"text": b.data, "bullet": "none"})
            elif b.kind == "sub":
                paras.append({"text": b.data, "bullet": "none", "bold": True, "color": "accent", "scale": 1.05})
            elif b.kind == "quote":
                paras.append({"text": f"「{b.data}」", "bullet": "none", "italic": True, "color": "muted"})
        return paras

    def render_content(self, spec):
        text_blocks = [b for b in spec.blocks if b.kind in TEXT_KINDS]
        visuals = [b for b in spec.blocks if b.kind not in TEXT_KINDS]
        paras = self.block_paras(text_blocks)

        if not visuals:
            s = self.new_slide("content")
            self.set_title(s, spec.title, "content")
            bodies = self.placeholders(s, BODY_TYPES)
            if bodies and paras:
                self.fill_text(bodies[0], paras, spec.title)
        elif not paras:
            s = self.new_slide("content", "title_only")
            self.set_title(s, spec.title, "content")
            self.place_visuals(s, visuals, self.content_area(), spec.title)
        else:
            use_two = "two_content" in self.layouts and len(
                [p for p in self.layouts["two_content"].placeholders if p.placeholder_format.type in BODY_TYPES]) >= 2
            s = self.new_slide("content", "two_content" if use_two else "content")
            self.set_title(s, spec.title, "content")
            bodies = self.placeholders(s, BODY_TYPES)
            if use_two and len(bodies) >= 2:
                left, right = bodies[0], bodies[1]
                vbox = box(right)
                right._element.getparent().remove(right._element)
            else:
                left = bodies[0] if bodies else None
                x, y, w, h = self.content_area() if left is None else box(left)
                ratio = float(self.cfg.get("text_ratio", 0.45))
                gap = int(Inches(0.3))
                tw = int(w * ratio)
                vbox = (x + tw + gap, y, w - tw - gap, h)
                if left is not None:
                    left.left, left.top, left.width, left.height = x, y, tw, h
            if left is not None:
                self.fill_text(left, paras, spec.title)
            self.place_visuals(s, visuals, vbox, spec.title)
        self.notes(s, spec.notes)
        self.cleanup(s)

    def fill_text(self, ph, paras, title):
        size, ok = fit_size(paras, ph.width, ph.height, self.sizes["body_max"], self.sizes["body_min"])
        if not ok:
            self.warnings.append(f"[{title}] text may overflow (shrunk to {size}pt) — split or trim")
        self.fill_paras(ph.text_frame, paras, "content", size)
        ph.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    def place_visuals(self, slide, visuals, area, title):
        x, y, w, h = area
        n = len(visuals)
        gap = int(Inches(0.25))
        if all(v.kind == "image" for v in visuals):
            cw = (w - gap * (n - 1)) // n
            cells = [(x + k * (cw + gap), y, cw, h) for k in range(n)]
        else:
            ch = (h - gap * (n - 1)) // n
            cells = [(x, y + k * (ch + gap), w, ch) for k in range(n)]
        for v, cell in zip(visuals, cells):
            getattr(self, f"place_{v.kind}")(slide, v.data, cell, title)

    def place_image(self, slide, d, cell, title):
        path = (self.base / d["path"]).resolve() if not Path(d["path"]).is_absolute() else Path(d["path"])
        x, y, w, h = cell
        if not path.exists():
            self.warnings.append(f"[{title}] image not found: {d['path']}")
            return
        if path.suffix.lower() == ".svg":
            self.warnings.append(f"[{title}] SVG is not supported, convert to PNG: {d['path']}")
            return
        cap_h = int(Inches(0.4)) if d["alt"] and self.cfg.get("image_captions", True) else 0
        from PIL import Image
        with Image.open(path) as im:
            iw, ih = im.size
        avail_h = h - cap_h
        scale = min(w / iw, avail_h / ih)
        pw, ph_ = int(iw * scale), int(ih * scale)
        px, py = x + (w - pw) // 2, y + (avail_h - ph_) // 2
        pic = slide.shapes.add_picture(str(path), px, py, pw, ph_)
        pic._element.nvPicPr.cNvPr.set("descr", d["alt"] or path.stem)
        if cap_h:
            tb = slide.shapes.add_textbox(x, py + ph_ + int(Inches(0.05)), w, cap_h)
            self.fill_paras(tb.text_frame, [{"text": d["alt"], "bullet": "none", "color": "muted",
                                             "align": PP_ALIGN.CENTER}], "content", self.sizes["caption"])

    def place_table(self, slide, d, cell, title):
        rows = d["rows"]
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        x, y, w, h = cell
        size = max(self.sizes["table_min"], min(self.sizes["table_max"], int(28 - 1.2 * len(rows) - ncols)))
        row_h = int(Pt(size * 2.1))
        shape = slide.shapes.add_table(len(rows), ncols, x, y, w, min(h, row_h * len(rows)))
        tbl = shape.table
        widths = [max(disp_width(plain(r[c])) for r in rows) + 2 for c in range(ncols)]
        tot = sum(widths)
        for c in range(ncols):
            tbl.columns[c].width = int(w * widths[c] / tot)
        aligns = d.get("aligns", [])
        amap = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
        for r, row in enumerate(rows):
            tbl.rows[r].height = row_h
            for c, val in enumerate(row):
                tf = tbl.cell(r, c).text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                if c < len(aligns):
                    p.alignment = amap[aligns[c]]
                fam = self.fonts.get("body")
                ea = self.fonts.get("ea_body") or self.fonts.get("ea")
                for chunk, fmt in inline_runs(val):
                    if chunk:
                        run = p.add_run()
                        run.text = chunk
                        set_font(run, family=fam, ea=ea, size=size, bold=True if (r == 0 or fmt.get("bold")) else None)
        est = row_h * len(rows)
        if est > h * 1.1:
            self.warnings.append(f"[{title}] table may overflow ({len(rows)} rows)")

    def place_code(self, slide, d, cell, title):
        x, y, w, h = cell
        lines = d["text"].split("\n")
        longest = max((len(l) for l in lines), default=1)
        size = self.sizes["code_max"]
        while size > self.sizes["code_min"] and (
                longest * size * 0.6 > (w / EMU_PER_PT - 22) or len(lines) * size * 1.25 > h / EMU_PER_PT - 22):
            size -= 1
        tb = slide.shapes.add_textbox(x, y, w, min(h, int(Pt(len(lines) * size * 1.25 + 26))))
        fill = tb.fill
        fill.solid()
        fill.fore_color.rgb = hex_rgb(self.colors.get("code_bg", "F3F4F6"))
        tf = tb.text_frame
        tf.word_wrap = True
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Inches(0.15))
        for k, ln in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = ln if ln else " "
            set_font(r, family=self.fonts.get("code", "Consolas"), size=size,
                     color=self.colors.get("code_text", "1F2937"))
        if len(lines) * size * 1.25 > h / EMU_PER_PT:
            self.warnings.append(f"[{title}] code block may overflow ({len(lines)} lines)")


def build_template(doc, tpl_path: Path | None, cfg: dict, theme: dict | None, base: Path):
    """Render `doc` onto a template (or python-pptx's blank master when tpl_path is None)."""
    prs = open_presentation(tpl_path)
    if tpl_path is None:
        scale_to_widescreen(prs)
    else:
        remove_all_slides(prs)
    keys = ("max_lines", "max_table_rows", "closing", "agenda", "agenda_title", "section_label")
    specs = plan(doc, {**cfg, **{k: v for k, v in doc.meta.items() if k in keys}})
    r = Renderer(prs, cfg, theme, base)
    for spec in specs:
        for piece in r.split_to_fit(spec):
            r.render(piece)
    return prs, r


def inspect_template(path: Path, write_config: Path | None = None) -> str:
    prs = open_presentation(path)
    W, H = prs.slide_width, prs.slide_height
    out = [f"slide size: {W / 914400:.2f}in × {H / 914400:.2f}in   masters: {len(prs.slide_masters)}   "
           f"existing slides: {len(prs.slides)}"]
    for i, l in enumerate(prs.slide_layouts):
        phs = []
        for ph in l.placeholders:
            t = str(ph.placeholder_format.type).split(".")[-1].split(" ")[0]
            if ph.placeholder_format.type in META_TYPES:
                continue
            geo = f"@{(ph.left or 0) / 914400:.1f},{(ph.top or 0) / 914400:.1f} {(ph.width or 0) / 914400:.1f}×{(ph.height or 0) / 914400:.1f}in"
            phs.append(f"{t}[{ph.placeholder_format.idx}]{geo}")
        out.append(f"[{i:2d}] {l.name}\n      " + ("  ".join(phs) or "(no placeholders)"))
    mapping = {k: v.name for k, v in detect_layouts(prs, {}).items()}
    out.append("\nauto-detected roles: " + json.dumps(mapping, ensure_ascii=False, indent=2))
    if write_config:
        cfg = {"file": path.name, "layouts": mapping,
               "sizes": {"body_max": 24, "body_min": 14}, "theme": {"fonts": {}}}
        Path(write_config).write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
        out.append(f"\nconfig written → {write_config}")
    return "\n".join(out)
