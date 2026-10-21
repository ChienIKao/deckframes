"""Canvas engine: draws every slide from design tokens (no template placeholders).

Used when the theme JSON has `"engine": "canvas"` (e.g. themes/blockframe.json).
Structure:  # cover · ## chapter · ### subsection · #### slide
Generated:  cover → outline (大綱) → per chapter: divider with sub-TOC → content slides
            (each with a top nav bar showing current chapter + subsection) → closing
"""
from __future__ import annotations

import math
import re
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_PATTERN
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from ..icons import IconError, add_svg_picture, parse_ref, recolor, svg_aspect
from .chrome import ChromeMixin
from .diagrams import DiagramMixin
from .style import StyleMixin
from .template import scale_to_widescreen
from ..icons import resolve as resolve_icon
from .style import mix
from ..layout import TEXT_KINDS, estimate_height_pt, merge_bullets, text_units
from ..markdown import Block, disp_width, inline_runs, is_cjk, plain

IN = 914400
W, H = 13.333, 7.5
PAD = 0.55
NAV_H = 1.0
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]


def E(x: float) -> int:
    return int(round(x * IN))


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


@dataclass
class Page:
    kind: str                      # cover | outline | divider | content | closing
    title: str = ""
    subtitle: str = ""
    chapter: int | None = None     # index into chapters
    sub: int | None = None         # index into chapter's subsections
    blocks: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    recap: bool = False


@dataclass
class View:
    """What one slide of a build sequence shows (see `+++` in deckframes-core)."""
    step: int = 10 ** 6            # blocks of steps <= this are drawn
    fade_before: int = -1          # blocks of steps < this are greyed out ...
    focus: str | None = None       # ... except these items (`+++ focus=2`)
    items: dict = field(default_factory=dict)  # id(component) -> (visible items, focused item, dim items before)


def parse_focus(spec):
    """`2,3` / `交叉注意力` → ({1, 2}, {names}) — numbers are 1-based item / row positions."""
    idx, names = set(), set()
    for tok in (spec or "").split(","):
        tok = tok.strip()
        if tok.isdigit():
            idx.add(int(tok) - 1)
        elif tok:
            names.add(tok)
    return idx, names


def in_focus(tags, idx, names):
    return tags.get("item") in idx or (tags.get("name") is not None and tags["name"] in names)


# --------------------------------------------------------------------------- XML helpers

RPR_AFTER_HL = ("a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill", "a:latin", "a:ea", "a:cs", "a:sym",
                "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")


def _insert_before(parent, el, tags):
    for t in tags:
        ref = parent.find(qn(t))
        if ref is not None:
            ref.addprevious(el)
            return
    parent.append(el)


def set_ea(rPr, face):
    el = rPr.find(qn("a:ea"))
    if el is None:
        el = etree.Element(qn("a:ea"))
        latin = rPr.find(qn("a:latin"))
        if latin is not None:
            latin.addnext(el)
        else:
            _insert_before(rPr, el, ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"))
    el.set("typeface", face)


def set_highlight(rPr, color):
    for old in rPr.findall(qn("a:highlight")):
        rPr.remove(old)
    hl = etree.Element(qn("a:highlight"))
    etree.SubElement(hl, qn("a:srgbClr"), val=color)
    _insert_before(rPr, hl, RPR_AFTER_HL)


def cjk_breaks(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("eaLnBrk", "1")
    pPr.set("hangingPunct", "1")
    pPr.set("latinLnBrk", "0")
    return pPr


def strip_style(shape):
    """Drop the theme style reference so no soft shadow / theme line sneaks in."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)
    return shape


# --------------------------------------------------------------------------- engine


class Canvas(StyleMixin, ChromeMixin, DiagramMixin):
    def __init__(self, theme: dict, meta: dict, base_dir: Path):
        self.t, self.meta, self.base = theme, meta, base_dir
        c = theme["colors"]
        self.black, self.white, self.ground = c["black"], c["white"], c["ground"]
        self.muted, self.ink = c["muted"], c["text"]
        self.palette = c["palette"]
        self.chapter_cycle = c.get("chapter_cycle", self.palette)
        f = theme["fonts"]
        self.f_display, self.f_label, self.f_body = f["display"], f["label"], f["body"]
        self.f_ea, self.f_code = f["ea"], f.get("code", "Consolas")
        s = theme.get("stroke", {})
        self.bw, self.sw = s.get("border", 2.25), s.get("shadow", 5)
        self.tw, self.tsw = s.get("thin", 1.5), s.get("thin_shadow", 3)
        self.sizes = {"title": 28, "subtitle": 16, "body_max": 20, "body_min": 12, "split_below": 14,
                      **theme.get("sizes", {})}
        self.prs = Presentation()
        scale_to_widescreen(self.prs)
        self.blank = self.prs.slide_layouts[6]
        # highlighter: prefer the 4th accent (yellow in BlockFrame) but only if ink stays readable on it
        light = [p for p in self.palette[3:] + self.palette[:3] if self.luminance(p) >= 0.55]
        self.hl = c.get("highlight") or (light[0] if light else "FFE58A")
        self.deco = theme.get("decorations", True)
        self.tilt = theme.get("tilt", True)
        # colour roles: ink on the slide ground, ink on cards, ink on light/dark fills
        self.title_ink = c.get("title", self.ink)
        self.dark_ink = c.get("on_light") or (self.black if self.luminance(self.black) < 0.4 else "111111")
        self.light_ink = c.get("on_dark", "FFFFFF")
        self.card_ink = c.get("card_text") or self.on(self.white)
        self.card_muted = c.get("card_muted") or ("4A4A4A" if self.luminance(self.white) >= 0.5 else "C9C9D3")
        self.tones = {"negative": c.get("negative", "E5534B"), "positive": c.get("positive", "4CAF6A"),
                      "info": c.get("info", self.palette[min(1, len(self.palette) - 1)])}
        self.init_style(theme)
        self.warnings: list[str] = []
        self._img_cache: dict = {}
        self._marks: list = []
        self.chapters = []
        self.counter = 0

    # ------------------------------------------------------------------ primitives
    @staticmethod
    def luminance(h: str) -> float:
        r, g, b = (int(h.lstrip("#")[i:i + 2], 16) / 255 for i in (0, 2, 4))
        return 0.2126 * r + 0.7152 * g + 0.0722 * b

    def on(self, fill: str) -> str:
        """Readable text colour on top of `fill`."""
        return getattr(self, "dark_ink", "111111") if self.luminance(fill) >= 0.5 else getattr(self, "light_ink", "FFFFFF")

    def dim_on(self, fill: str) -> str:
        return "4A4A4A" if self.luminance(fill) >= 0.5 else "D8D8D8"

    def shape(self, s, kind, x, y, w, h, fill=None, line=None, lw=None, rot=0.0, pattern=None):
        sp = strip_style(s.shapes.add_shape(kind, E(x), E(y), E(w), E(h)))
        if pattern:
            sp.fill.patterned()
            sp.fill.pattern = pattern
            sp.fill.fore_color.rgb = rgb(fill[0])
            sp.fill.back_color.rgb = rgb(fill[1])
        elif fill:
            sp.fill.solid()
            sp.fill.fore_color.rgb = rgb(fill)
        else:
            sp.fill.background()
        width = self.bw if lw is None else lw
        if line and width > 0:
            sp.line.color.rgb = rgb(line)
            sp.line.width = Pt(width)
        else:
            sp.line.fill.background()
        if rot and self.tilt:
            sp.rotation = rot
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = 0.5
        return sp

    def block(self, s, x, y, w, h, fill, thin=False, rot=0.0, kind=MSO_SHAPE.RECTANGLE, line=None,
              shadow=None):
        """A card / box drawn in the theme's surface style (brutal, glass, clay, neu, paper, …)."""
        return self.surface_block(s, x, y, w, h, fill, thin, rot, kind, line, shadow)

    def line(self, s, x1, y1, x2, y2, color=None, width=None):
        ln = strip_style(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2)))
        ln.line.color.rgb = rgb(color or self.black)
        ln.line.width = Pt(width or self.bw)
        return ln

    def runs(self, p, text, size, font="body", bold=None, color=None, italic=None, spc=None, hl_color=None):
        latin = {"display": self.f_display, "label": self.f_label, "body": self.f_body, "heading": self.f_heading,
                 "code": self.f_code}[font]
        ea_face = self.f_ea_heading if font == "heading" else self.f_ea
        for chunk, fmt in inline_runs(text):
            if not chunk:
                continue
            r = p.add_run()
            r.text = chunk
            f = r.font
            f.size = Pt(size)
            f.bold = True if fmt.get("bold") else bold
            if fmt.get("italic") or italic:
                f.italic = True
            f.color.rgb = rgb(color or self.ink)
            f.name = self.f_code if fmt.get("code") else latin
            rPr = r._r.get_or_add_rPr()
            set_ea(rPr, ea_face)
            if any(is_cjk(ch) for ch in chunk):
                rPr.set("lang", "zh-TW")
                rPr.set("altLang", "en-US")
            if spc:
                rPr.set("spc", str(int(spc * 100)))
            if fmt.get("hl"):
                set_highlight(rPr, (hl_color or self.hl).lstrip("#"))
            if fmt.get("link"):
                r.hyperlink.address = fmt["link"]

    def tf_setup(self, tf, anchor=MSO_ANCHOR.TOP, margin=0.0, wrap=True):
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Emu(E(margin)))

    def text(self, s, x, y, w, h, lines, size, font="body", bold=None, color=None, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spc=None, line_spacing=None, target=None, margin=0.0):
        """lines: str or list[str|dict(text,size,bold,color,font,space_after)]."""
        if target is None:
            target = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
        tf = target.text_frame
        self.tf_setup(tf, anchor, margin)
        if isinstance(lines, str):
            lines = [lines]
        for k, ln in enumerate(lines):
            d = ln if isinstance(ln, dict) else {"text": ln}
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = d.get("align", align)
            cjk_breaks(p)
            if line_spacing:
                p.line_spacing = line_spacing
            if d.get("space_after") is not None:
                p.space_after = Pt(d["space_after"])
            self.runs(p, d["text"], d.get("size", size), d.get("font", font), d.get("bold", bold),
                      d.get("color", color), spc=d.get("spc", spc))
        return target

    def measure(self, text, size, font="body"):
        """Approximate rendered width in inches."""
        factor = {"display": 1.18, "label": 1.0, "body": 1.0, "heading": 1.05, "code": 1.0}[font]
        w = 0.0
        for ch in plain(text):
            w += disp_width(ch) * (factor if ord(ch) < 0x2E80 else 1.0)
        return w * size / 72

    def pill(self, s, x, y, text, fill, size=11, color=None, rot=0.0, h=None, shadow=True, upper=True,
             font="label", line=None):
        text = text.upper() if upper and text.isascii() else text
        h = h or size / 72 * 2.0
        w = self.pill_width(text, size, font, upper)
        if shadow:
            sp = self.block(s, x, y, w, h, fill, thin=True, rot=rot, kind=MSO_SHAPE.ROUNDED_RECTANGLE, line=line)
        else:
            sp = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line or self.black,
                            lw=self.tw, rot=rot)
        self.text(s, 0, 0, 0, 0, text, size, font=font, bold=True, color=color or self.on(fill),
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spc=1 if upper else None, target=sp)
        return sp, w

    def lines_height(self, lines, width):
        """Estimated height (in) of text() lines inside `width` inches."""
        total = 0.0
        for ln in lines:
            d = ln if isinstance(ln, dict) else {"text": ln}
            sz = d.get("size", 14)
            n = max(1, math.ceil(self.measure(d["text"], sz) / max(0.5, width)))
            total += n * sz * 1.25 / 72 + d.get("space_after", 0) / 72
        return total

    def card_lines(self, it, title=20, desc=15, child=14):
        lines = [{"text": it["title"], "size": title, "bold": True, "color": self.card_ink, "space_after": 3}]
        if it["desc"]:
            lines.append({"text": it["desc"], "size": desc, "color": self.card_muted, "space_after": 2})
        for c in it["children"] + it["extra"]:
            lines.append({"text": "■ " + c, "size": child, "color": self.card_muted})
        return lines

    def pill_width(self, text, size, font="label", upper=True):
        return self.measure(text, size, font) + 0.36 + (len(text) * size * 0.08 / 72 if upper else 0)

    def star(self, s, x, y, d, fill, text=None, rot=12, size=14):
        if not self.deco:
            return None
        sp = self.shape(s, MSO_SHAPE.STAR_10_POINT, x, y, d, d, fill=fill, line=self.black, lw=self.tw, rot=rot)
        if text:
            self.text(s, 0, 0, 0, 0, text, size, font="display", color=self.on(fill), align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE, target=sp)
        return sp

    def stripes(self, s, x, y, w, h, color, rot=0.0):
        if not self.deco:
            return None
        return self.shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=(self.black, color), line=self.black,
                          lw=self.tw, rot=rot, pattern=MSO_PATTERN.WIDE_UPWARD_DIAGONAL)

    def dots(self, s, x, y, w, h):
        if not self.deco:
            return None
        return self.shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=(self.muted, self.ground),
                          pattern=MSO_PATTERN.PERCENT_5)

    def picture(self, s, path, x, y, w, h):
        from PIL import Image
        p = Path(path)
        p = p if p.is_absolute() else (self.base / p)
        if not p.exists():
            self.warnings.append(f"image not found: {path}")
            return None
        if p.suffix.lower() == ".svg":
            data = p.read_bytes()
            iw, ih = svg_aspect(data), 1.0
        else:
            with Image.open(p) as im:
                iw, ih = im.size
        sc = min(w / iw, h / ih)
        pw, ph = iw * sc, ih * sc
        box = (E(x + (w - pw) / 2), E(y + (h - ph) / 2), E(pw), E(ph))
        if p.suffix.lower() == ".svg":
            pic = add_svg_picture(s, data, *box, descr=p.stem)
        else:
            pic = s.shapes.add_picture(str(p), *box)
            pic._element.nvPicPr.cNvPr.set("descr", p.stem)
        return pic, (x + (w - pw) / 2, y + (h - ph) / 2, pw, ph)

    def glyph(self, s, ref, x, y, w, h, color, scale=0.6) -> bool:
        """Draw an icon reference centred in the box. False when `ref` is plain text (caller draws it)."""
        if not ref or parse_ref(ref)[0] == "text":
            return False
        try:
            kind, data, recolourable = resolve_icon(ref, self.base)
        except IconError as e:
            self.warnings.append(str(e))
            return False
        side = min(w, h) * scale
        if kind == "svg":
            if recolourable:
                data = recolor(data, color)
            a = svg_aspect(data)
            gw, gh = (side, side / a) if a >= 1 else (side * a, side)
            add_svg_picture(s, data, E(x + (w - gw) / 2), E(y + (h - gh) / 2), E(gw), E(gh), descr=ref)
        else:
            import io
            from PIL import Image
            with Image.open(io.BytesIO(data)) as im:
                a = im.size[0] / im.size[1]
            gw, gh = (side, side / a) if a >= 1 else (side * a, side)
            s.shapes.add_picture(io.BytesIO(data), E(x + (w - gw) / 2), E(y + (h - gh) / 2), E(gw), E(gh))
        return True

    def pal(self, k):
        return self.palette[k % len(self.palette)]

    def color_for(self, k):
        return self.palette[k % len(self.palette)]

    def ch_color(self, ci):
        return self.chapter_cycle[ci % len(self.chapter_cycle)]

    def new_slide(self, ground=None, notes=None, role=None):
        s = self.prs.slides.add_slide(self.blank)
        self._marks = []
        self.paint_ground(s, ground, role)
        self.counter += 1
        if notes:
            s.notes_slide.notes_text_frame.text = "\n\n".join(notes)
        return s

    # ------------------------------------------------------------------ shape tagging (build steps / focus)
    @contextmanager
    def tag(self, s, **kw):
        """Every shape drawn inside the block gets these tags (inner tags win), e.g. item=k for the k-th item
        of a component, so later passes can hide or grey out single items."""
        tree = s.shapes._spTree
        start = len(tree)
        try:
            yield
        finally:
            self._marks.append((start, len(tree), kw))

    def tagged(self, s, start):
        """[(element, tags)] for the shapes drawn since spTree position `start`."""
        kids = list(s.shapes._spTree)
        marks = [m for m in self._marks if m[0] >= start]
        out = []
        for i in range(start, len(kids)):
            tags = {}
            for _, _, kw in sorted((m for m in marks if m[0] <= i < m[1]), key=lambda m: m[0] - m[1]):
                tags.update(kw)
            out.append((kids[i], tags))
        return out

    @staticmethod
    def el_box(el):
        """(x, y, w, h) in inches of a slide element, or None."""
        xfrm = el.find(qn("p:spPr") + "/" + qn("a:xfrm"))
        if xfrm is None:
            xfrm = el.find(qn("p:xfrm"))
        if xfrm is None:
            xfrm = el.find(qn("p:grpSpPr") + "/" + qn("a:xfrm"))
        if xfrm is None:
            return None
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        if off is None or ext is None:
            return None
        return (int(off.get("x")) / IN, int(off.get("y")) / IN, int(ext.get("cx")) / IN, int(ext.get("cy")) / IN)

    @staticmethod
    def union(boxes):
        boxes = [b for b in boxes if b]
        if not boxes:
            return None
        x0, y0 = min(b[0] for b in boxes), min(b[1] for b in boxes)
        x1, y1 = max(b[0] + b[2] for b in boxes), max(b[1] + b[3] for b in boxes)
        return x0, y0, x1 - x0, y1 - y0

    # ------------------------------------------------------------------ structure
    def build(self, doc):
        meta = self.meta
        self.chapters = [c for c in doc.sections if c.title]
        pages = [Page("cover", title=doc.cover.title, subtitle=doc.cover.subtitle or meta.get("subtitle", ""),
                      notes=doc.cover.notes)]
        if doc.cover.blocks:
            pages.append(Page("content", title=doc.cover.title, blocks=doc.cover.blocks))
        if meta.get("outline", True) and len(self.chapters) >= 2:
            pages.append(Page("outline"))
        recap = bool(meta.get("recap", False))
        ci = -1
        for sec in doc.sections:
            if sec.title:
                ci += 1
                pages.append(Page("divider", chapter=ci, notes=sec.notes))
                intro = [b for b in sec.blocks]
                if intro and not (not sec.subsections and self.short_intro(intro)):
                    pages.append(Page("content", title=sec.title, chapter=ci, blocks=intro))
            elif sec.blocks:
                pages.append(Page("content", blocks=sec.blocks, notes=sec.notes))
            chap = ci if sec.title else None
            for sl in sec.slides:
                pages.append(Page("content", title=sl.title, subtitle=sl.subtitle, chapter=chap,
                                  blocks=sl.blocks, notes=sl.notes))
            for si, sub in enumerate(sec.subsections):
                if recap and si > 0 and chap is not None:
                    pages.append(Page("divider", chapter=chap, sub=si, recap=True))
                if sub.blocks or not sub.slides:
                    pages.append(Page("content", title=sub.title, subtitle=sub.subtitle, chapter=chap, sub=si,
                                      blocks=sub.blocks, notes=sub.notes))
                for sl in sub.slides:
                    pages.append(Page("content", title=sl.title, subtitle=sl.subtitle, chapter=chap, sub=si,
                                      blocks=sl.blocks, notes=sl.notes))
        closing = meta.get("closing", "Q&A")
        if closing:
            pages.append(Page("closing", title=str(closing) if closing is not True else "Q&A"))

        for pg in pages:
            if pg.kind == "content" and self.stepped(pg):
                self.render_stepped(pg)
            elif pg.kind == "content":
                for piece in self.paginate(pg):
                    self.render_content(piece)
            else:
                getattr(self, f"render_{pg.kind}")(pg)
        return self.prs

    @staticmethod
    def short_intro(blocks):
        return len(blocks) == 1 and blocks[0].kind in ("para", "quote") and disp_width(plain(blocks[0].data)) <= 80

    # ------------------------------------------------------------------ cover
    def cover_split(self, pg):
        s = self.new_slide(notes=pg.notes, role="cover")
        m = self.meta
        self.dots(s, 7.4, 0, 5.93, 3.6)
        _, pw = self.pill(s, PAD + 0.1, 1.0, m.get("eyebrow", "PRESENTATION"), self.pal(3), size=13)
        title = pg.title
        # largest size (≤54pt) that splits the title into 1–3 evenly filled lines
        n = max(1.0, disp_width(title))
        size = max(min(54, int(7.1 * 72 / math.ceil(n / k) * 0.97)) for k in (1, 2, 3))
        lines = math.ceil(n * size / 72 / 7.1)
        tw = min(7.3, math.ceil(n / lines) * size / 72 * 1.04 + 0.15)   # narrow the box so lines balance
        self.text(s, PAD + 0.1, 1.75, tw, 3.2, title, size, bold=True, color=self.title_ink, font="heading",
                  anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
        if pg.subtitle:
            sub_font = "display" if pg.subtitle.isascii() else "body"
            sub = pg.subtitle.upper() if pg.subtitle.isascii() else pg.subtitle
            self.text(s, PAD + 0.1, 5.0, 7.3, 0.6, sub, 18, font=sub_font, bold=True, color=self.title_ink)
        x = PAD + 0.1
        for k, key in enumerate(("author", "date")):
            if m.get(key):
                _, w = self.pill(s, x, 6.25, str(m[key]), self.white, size=12, upper=False, font="body")
                x += w + 0.25
        # right-hand decorations
        if m.get("cover_image"):
            self.picture_card(s, m["cover_image"], 8.3, 1.3, 4.3, 3.6, rot=3)
        else:
            self.decorate(s, (7.9, 0.5, 4.9, 6.2), "cover")

    def picture_card(self, s, path, x, y, w, h, rot=0.0):
        got = self.picture(s, path, x, y, w, h)
        if got:
            self.outline(got[0])
            if rot:
                got[0].rotation = rot
        return got

    def outline(self, pic):
        """Images get only a thin ink outline (theme `stroke.image`, pt) — no card, no padding, no shadow."""
        width = self.t.get("stroke", {}).get("image", 1.0)
        if width > 0:
            pic.line.color.rgb = rgb(self.black)
            pic.line.width = Pt(width)

    # ------------------------------------------------------------------ outline (大綱)
    def outline_line(self, pg):
        s = self.new_slide(role="outline")
        m = self.meta
        self.pill(s, PAD + 0.1, 1.0, m.get("outline_label", "OUTLINE"), self.pal(3), size=13)
        self.text(s, PAD + 0.1, 1.7, 5.8, 1.3, m.get("outline_title", "大綱"), 60, bold=True, color=self.title_ink,
                  font="heading")
        self.text(s, PAD + 0.1, 2.9, 5.8, 0.7, m.get("outline_en", "OUTLINE"), 26, font="display",
                  color=self.title_ink, spc=-0.5)
        self.decorate(s, (0.8, 4.0, 4.9, 2.9), "outline")

        n = len(self.chapters)
        top, bottom = 0.7, 6.8
        step = min(1.05, (bottom - top) / n)
        y0 = top + ((bottom - top) - step * n) / 2
        lx = 7.3
        self.line(s, lx, 0, lx, H, width=self.rule_w)
        bh = min(0.62, step * 0.62)
        for k, ch in enumerate(self.chapters):
            cy = y0 + step * k + step / 2
            self.block(s, lx - 0.48, cy - bh / 2, 0.96, bh, self.ch_color(k))
            self.text(s, lx - 0.48, cy - bh / 2, 0.96, bh, f"{k + 1:02d}", 20, font="display",
                      color=self.on(self.ch_color(k)), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            lines = [{"text": ch.title, "size": 24, "bold": True, "color": self.title_ink}]
            if ch.subtitle:
                lines.append({"text": ch.subtitle.upper(), "size": 11, "font": "label", "bold": True,
                              "color": self.muted, "spc": 1})
            self.text(s, lx + 0.9, cy - step / 2, 4.9, step, lines, 24, anchor=MSO_ANCHOR.MIDDLE)
        self.counter_pill(s)

    # ------------------------------------------------------------------ chapter divider + sub-TOC
    def divider_panel(self, pg):
        ch = self.chapters[pg.chapter]
        col = self.ch_color(pg.chapter)
        s = self.new_slide(notes=pg.notes, role="divider")
        pw = 5.3
        self.shape(s, MSO_SHAPE.RECTANGLE, -0.05, -0.05, pw + 0.05, H + 0.1, fill=col,
                   line=self.black if self.surface in ("brutal", "pixel", "outline") else None)
        self.text(s, PAD + 0.05, 0.45, 3.5, 1.4, f"{pg.chapter + 1:02d}", 80, font="display", color=self.on(col),
                  spc=-2)
        self.shape(s, MSO_SHAPE.RECTANGLE, PAD + 0.15, 1.85, 1.9, 0.1, fill=self.on(col))
        n = disp_width(ch.title)
        size = 60 if n <= 4 else 50 if n <= 6 else 40 if n <= 9 else 32
        lines = [{"text": ch.title, "size": size, "bold": True, "color": self.on(col), "space_after": 6,
                  "font": "heading"}]
        if ch.subtitle:
            lines.append({"text": ch.subtitle.upper(), "size": 22, "font": "display", "color": self.on(col),
                          "spc": -0.3})
        self.text(s, 0.3, 2.5, pw - 0.6, 3.0, lines, size, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        self.decorate(s, (pw - 1.2, 5.4, 1.8, 1.8), "divider")

        subs = ch.subsections
        lx = 6.6
        if not subs:
            intro = ch.blocks[0].data if ch.blocks and self.short_intro(ch.blocks) else ""
            if intro:
                self.block(s, 6.4, 2.6, 6.2, 2.2, self.white)
                self.text(s, 6.7, 2.8, 5.6, 1.8, intro, 22, bold=True, color=self.card_ink,
                          anchor=MSO_ANCHOR.MIDDLE)
            self.counter_pill(s)
            return
        self.line(s, lx, 0, lx, H, width=self.rule_w)
        n = len(subs)
        top, bottom = 0.6, 6.9
        step = min(1.15, (bottom - top) / n)
        y0 = top + ((bottom - top) - step * n) / 2
        bh = min(0.66, step * 0.6)
        for k, sub in enumerate(subs):
            cy = y0 + step * k + step / 2
            current = pg.recap and k == pg.sub
            dim = pg.recap and k != pg.sub
            fill = self.title_ink if current else (self.white if dim else col)
            self.block(s, lx - 0.45, cy - bh / 2, 0.9, bh, fill, thin=dim)
            self.text(s, lx - 0.45, cy - bh / 2, 0.9, bh, ROMAN[k] if k < len(ROMAN) else str(k + 1), 20,
                      font="display", color=self.muted if dim else self.on(fill),
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            lines = [{"text": sub.title, "size": 24, "bold": True,
                      "color": self.muted if dim else self.title_ink}]
            if sub.subtitle:
                lines.append({"text": sub.subtitle, "size": 12, "color": self.muted})
            self.text(s, lx + 0.8, cy - step / 2, 5.6, step, lines, 24, anchor=MSO_ANCHOR.MIDDLE)
        self.counter_pill(s)

    # ------------------------------------------------------------------ closing
    def closing_frame(self, pg):
        s = self.new_slide(ground=self.dark_ink, role="closing")
        w, h = 8.2, 2.8
        x, y = (W - w) / 2, (H - h) / 2 + 0.2
        self.shape(s, MSO_SHAPE.RECTANGLE, x + 0.17, y + 0.17, w, h, fill=self.pal(3))
        frame = self.shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=self.dark_ink, line="FFFFFF", lw=3)
        size = 72 if disp_width(pg.title) <= 6 else 48
        font = "display" if pg.title.isascii() else "body"
        self.text(s, 0, 0, 0, 0, pg.title, size, font=font, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, target=frame)
        label = self.meta.get("closing_label", "THANK YOU")
        self.pill(s, (W - self.pill_width(label, 13)) / 2, y - 0.75, label, "FFFFFF", size=13, shadow=False)
        self.star(s, x + w - 0.7, y - 0.6, 1.3, self.pal(0), rot=18)

    # ------------------------------------------------------------------ chrome for content slides
    def counter_pill(self, s):
        self.pill(s, W - PAD - 0.75, H - 0.52, f"{self.counter:02d}", self.white, size=11, shadow=False)

    def nav_band(self, s, ci, si):
        col = self.ch_color(ci)
        self.shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, NAV_H, fill=col)
        self.line(s, 0, NAV_H, W, NAV_H)
        left, right = PAD, W - PAD
        if self.meta.get("logo_left"):
            self.picture(s, self.meta["logo_left"], PAD, 0.1, 0.8, 0.8)
            left += 1.0
        if self.meta.get("logo_right"):
            self.picture(s, self.meta["logo_right"], W - PAD - 0.8, 0.1, 0.8, 0.8)
            right -= 1.0
        n = len(self.chapters)
        slot = (right - left) / n
        for k, ch in enumerate(self.chapters):
            cx = left + slot * k + slot / 2
            if k == ci:
                tw = self.measure(ch.title, 14) + 0.5
                sp = self.block(s, cx - tw / 2, 0.1, tw, 0.4, self.white, thin=True)
                self.text(s, 0, 0, 0, 0, ch.title, 14, bold=True, color=self.card_ink, align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE, target=sp)
            else:
                self.text(s, cx - slot / 2, 0.1, slot, 0.4, ch.title, 13, color=self.dim_on(col),
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        subs = self.chapters[ci].subsections
        if not subs:
            return
        size = 11
        widths = [self.measure(sub.title, size) + 0.4 for sub in subs]
        gap = 0.18
        total = sum(widths) + gap * (len(subs) - 1)
        if total > right - left:
            size = 9
            widths = [self.measure(sub.title, size) + 0.3 for sub in subs]
            total = sum(widths) + gap * (len(subs) - 1)
        cx = left + slot * ci + slot / 2
        x = min(max(left, cx - total / 2), right - total)
        for k, (sub, w) in enumerate(zip(subs, widths)):
            if k == si:
                sp = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 0.6, w, 0.3, fill=self.dark_ink,
                                line=self.dark_ink, lw=self.tw)
                self.text(s, 0, 0, 0, 0, sub.title, size, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE, target=sp)
            else:
                self.text(s, x, 0.6, w, 0.3, sub.title, size, color=self.dim_on(col), align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE)
            x += w + gap

    def full_title(self, pg):
        if pg.chapter is None or pg.sub is None:
            return pg.title
        sub = self.chapters[pg.chapter].subsections[pg.sub].title
        if pg.title == sub or pg.title.startswith(sub):
            return pg.title
        fmt = self.meta.get("title_format", "{sub} – {title}")
        return fmt.format(sub=sub, title=pg.title)

    # ------------------------------------------------------------------ pagination
    def split_blocks(self, blocks):
        text = [b for b in blocks if b.kind in TEXT_KINDS]
        visuals = [b for b in blocks if b.kind in ("image", "table", "code", "component")]
        callouts = [b for b in blocks if b.kind == "callout"]
        sources = [b for b in blocks if b.kind == "source"]
        return text, visuals, callouts, sources

    def text_ratio(self, visuals):
        types = {v.data["type"] for v in visuals if v.kind == "component"}
        if types & {"diagram", "lanes", "matrix", "mapping", "cycle"}:
            return 0.33
        if types & {"cards", "steps", "compare", "pyramid", "funnel", "stack"}:
            return 0.42
        return 0.45

    def paginate(self, pg):
        text, visuals, callouts, sources = self.split_blocks(pg.blocks)
        max_rows = int(self.meta.get("max_table_rows", 9))
        vis = []
        for v in visuals:
            if v.kind == "table" and len(v.data["rows"]) - 1 > max_rows:
                head, body = v.data["rows"][0], v.data["rows"][1:]
                for k in range(0, len(body), max_rows):
                    vis.append(Block("table", {"rows": [head] + body[k:k + max_rows], "aligns": v.data["aligns"]}))
            else:
                vis.append(v)
        first_vis, rest_vis = (vis[:1], vis[1:]) if text else (vis[:2], vis[2:])
        if len(first_vis) == 2 and not all(v.kind == "image" for v in first_vis):
            rest_vis = first_vis[1:] + rest_vis
            first_vis = first_vis[:1]
        pages = self.fit_text(pg, text, first_vis, callouts, sources)
        for v in rest_vis:
            pages.append(Page("content", title=pg.title, subtitle=pg.subtitle, chapter=pg.chapter, sub=pg.sub,
                              blocks=[v]))
        suffix = self.meta.get("continued_suffix", "（續）")
        for p in pages[1:]:
            if not p.title.endswith(suffix):
                p.title = pg.title + suffix
        return pages

    def fit_text(self, pg, text, vis, callouts, sources):
        base = Page("content", title=pg.title, subtitle=pg.subtitle, chapter=pg.chapter, sub=pg.sub,
                    notes=pg.notes)
        units = [u for b in text for u, _ in text_units(b)]
        floor = self.sizes["split_below"]
        paras = self.text_paras(text)
        lay = self.layout(pg, paras, vis, callouts, bool(sources))
        w, h = lay["text"][2], lay["text"][3]
        if len(units) < 2 or estimate_height_pt(paras, floor * 1.08, E(w)) <= h * 72:
            base.blocks = text + vis + callouts + sources
            return [base]
        mid = len(units) // 2
        a = self.fit_text(pg, merge_bullets(units[:mid]), vis, callouts, sources)
        b_pg = Page("content", title=pg.title, subtitle=pg.subtitle, chapter=pg.chapter, sub=pg.sub)
        b = self.fit_text(b_pg, merge_bullets(units[mid:]), [], [], [])
        return a + b

    # ------------------------------------------------------------------ content slide
    def text_paras(self, blocks):
        paras = []
        for b in blocks:
            if b.kind == "bullets":
                for it in b.data:
                    paras.append({"text": it["text"], "level": it["level"], "ordered": it["ordered"],
                                  "bullet": "sq", "scale": max(0.8, 1 - 0.1 * it["level"])})
            elif b.kind == "para":
                paras.append({"text": b.data, "bullet": "none"})
            elif b.kind == "sub":
                paras.append({"text": b.data, "bullet": "none", "bold": True, "scale": 1.05})
            elif b.kind == "quote":
                paras.append({"text": f"「{b.data}」", "bullet": "none", "italic": True, "color": self.muted})
        return paras

    # ------------------------------------------------------------------ layout planning
    def img_info(self, d):
        """(aspect w/h, pixel width or None) for an image block, cached."""
        key = d["path"]
        if key not in self._img_cache:
            p = Path(d["path"])
            p = p if p.is_absolute() else self.base / p
            info = None
            if p.exists():
                if p.suffix.lower() == ".svg":
                    info = (svg_aspect(p.read_bytes()), None)
                else:
                    from PIL import Image
                    with Image.open(p) as im:
                        info = (im.size[0] / im.size[1], im.size[0])
            self._img_cache[key] = info
        return self._img_cache[key]

    def wrapped_lines(self, paras, w, size):
        """Lines the paragraphs need in a box of width w at size (per-character widths, Latin-friendly)."""
        return sum(max(1, math.ceil(self.measure(pd["text"], size) / max(0.5, w - 0.4 - 0.3 * pd.get("level", 0))))
                   for pd in paras)

    def callout_height(self, d, w):
        need = self.lines_height([{"text": d["text"], "size": 15}], max(1.0, w - 1.6)) + 0.3
        return min(1.6, max(0.85, need))

    def layout(self, pg, paras, visuals, callouts, has_sources):
        """Decide boxes for this slide. Returns {mode, text, visual, callouts:[(x, y, w, h)], title_w}.

        Images are sized from their real aspect ratio instead of a fixed column split:
        - image only, portrait  → "showcase": image fills the full height on the right, the subtitle
                                  becomes the headline on the left
        - image only, landscape → image gets the whole body (down to the footer)
        - text + one image      → the image column is as wide as the image is at full height
                                  (clamped), text takes the rest; callouts sit under the text
        `"left"` / `"wide"` in the image's Markdown title override placement / emphasis.
        """
        top = self.content_top(pg)
        x, w = PAD, W - 2 * PAD
        bottom = H - (0.75 if has_sources else 0.68)
        body_y = top + 1.05 + (0.4 if pg.subtitle else 0)
        gap = 0.45
        single = visuals[0] if len(visuals) == 1 and visuals[0].kind == "image" else None
        info = self.img_info(single.data) if single else None
        out = {"mode": "", "title_w": w, "text": (x, body_y, w, bottom - body_y), "visual": None, "callouts": []}

        def stack_callouts(cx, cw, y_end):
            hs = [self.callout_height(c.data, cw) for c in callouts]
            y0 = y_end - sum(hs) - 0.3 * len(hs)
            out["callouts"] = []
            yy = y0 + 0.3
            for hh in hs:
                out["callouts"].append((cx, yy, cw, hh))
                yy += hh + 0.3
            return y0

        if single and info and not paras:
            a, opts = info[0], single.data.get("opts", [])
            cb = stack_callouts(x, w, bottom) if callouts else bottom

            def fitted(bw, bh):  # area of the image fitted into a box
                return min(bw, bh * a) * min(bh, bw / a)

            vy = top + 0.05
            sc_h = bottom - vy - 0.05
            sc_w = min(w * (0.5 if a < 1.0 else 1.0) - 0.05, sc_h * a, w - 3.2 - gap - 0.05)
            showcase_area = fitted(sc_w, sc_h)
            plain_area = fitted(w - 0.05, cb - body_y - 0.05)
            wants_showcase = (pg.subtitle or single.data.get("alt")) and "center" not in opts
            if wants_showcase and (a < 1.0 or showcase_area > plain_area * 1.1):
                img_w = sc_w + 0.05
                left = "left" in opts
                vx = x if left else x + w - img_w
                tx = x + img_w + gap if left else x
                tw = w - img_w - gap
                out.update(mode="showcase", title_w=tw, visual=(vx, vy, img_w, bottom - vy))
                tb = stack_callouts(tx, tw, bottom) if callouts else bottom
                out["text"] = (tx, body_y - (0.4 if pg.subtitle else 0), tw, tb - body_y + (0.4 if pg.subtitle else 0))
                if left:
                    out["title_x"] = tx
                return out
            vb = stack_callouts(x, w, bottom) if callouts else bottom
            out.update(mode="visual", visual=(x, body_y, w, vb - body_y))
            return out

        mode_below = str(self.meta.get("visual_below", "")).lower()
        opts0 = single.data.get("opts", []) if single else []
        bleed = float(self.meta.get("visual_bleed", 0) or 0)
        below = "below" in opts0 or (mode_below in ("true", "yes", "1") and "side" not in opts0)
        if visuals and paras and not below and mode_below == "auto" and "side" not in opts0:
            # try the side-by-side split; go below only when its text column would keep wrapping
            if single and info:
                cap = 0.5 if single.data.get("alt") else 0.0
                nat = (bottom - body_y - cap - 0.05) * info[0] + 0.05
                lo, hi = (0.55 * w, 0.75 * w) if "wide" in opts0 else (0.3 * w, 0.66 * w)
                side_tw = w - max(lo, min(hi, nat)) - gap
            else:
                side_tw = w * self.text_ratio(visuals)
            below = self.wrapped_lines(paras, side_tw, self.sizes["body_max"] - 2) - len(paras) > 1
        if visuals and paras and below:
            # text across the full width on top, visual underneath (long-sentence languages)
            # the visual may run down to the source line (or near the slide edge when there is none)
            low = H - (0.62 if has_sources else 0.25)
            vb = stack_callouts(x, w, low) if callouts else low
            # same estimate draw_text fits with, so the text keeps its body_max - 4 size
            th = estimate_height_pt(paras, (self.sizes["body_max"] - 4) * 1.08, E(w)) / 72 + 0.01  # + rounding slack
            th = min(th, 0.4 * (vb - body_y))
            ty = body_y - 0.25
            vy = ty + th + 0.05
            out.update(mode="below", text=(x, ty, w, th), visual=(x - bleed, vy, w + 2 * bleed, vb - vy))
            return out

        if single and info and paras:
            a, opts = info[0], single.data.get("opts", [])
            cap = 0.5 if single.data.get("alt") else 0.0
            full_h = bottom - body_y
            natural = (full_h - cap - 0.05) * a + 0.05
            lo, hi = (0.55 * w, 0.75 * w) if "wide" in opts else (0.3 * w, 0.66 * w)
            img_w = max(lo, min(hi, natural))
            tw = w - img_w - gap
            floor = self.sizes["split_below"]
            text_bottom = bottom
            for _ in range(12):  # give the text room if it cannot fit at the split threshold
                text_bottom = (stack_callouts(0, tw, bottom) if callouts else bottom)
                if tw >= 0.62 * w or estimate_height_pt(paras, floor * 1.08, E(tw)) <= (text_bottom - body_y) * 72:
                    break
                tw += 0.3
                img_w = w - tw - gap
            left = "left" in opts
            tx = x + img_w + gap if left else x
            vx = x if left else x + tw + gap
            if callouts:
                stack_callouts(tx, tw, bottom)
            vx, img_w = (vx - bleed, img_w + bleed) if left else (vx, img_w + bleed)
            out.update(mode="split", text=(tx, body_y, tw, text_bottom - body_y), visual=(vx, body_y, img_w, full_h))
            return out

        vb = stack_callouts(x, w, bottom) if callouts else bottom
        h = vb - body_y
        if not visuals:
            out.update(mode="text", text=(x, body_y, w, h))
        elif not paras:
            out.update(mode="visual", visual=(x, body_y, w, h))
        else:
            tw = w * self.text_ratio(visuals)
            out.update(mode="split", text=(x, body_y, tw, h), visual=(x + tw + gap, body_y, w - tw - gap + bleed, h))
        return out

    def content_title(self, s, pg, lay):
        """Slide title (+ grey subtitle line); returns the shapes so an overlay can lift them above its veil."""
        top = self.content_top(pg)
        out = [self.text(s, lay.get("title_x", PAD), top, lay["title_w"], 0.7, self.full_title(pg),
                         self.sizes["title"], bold=True, color=self.title_ink, anchor=MSO_ANCHOR.MIDDLE,
                         font="heading")]
        if pg.subtitle and lay["mode"] != "showcase":
            out.append(self.text(s, PAD, top + 0.7, W - 2 * PAD, 0.4, pg.subtitle, self.sizes["subtitle"], bold=True,
                                 color=self.muted))
        return out

    # ------------------------------------------------------------------ build steps (`+++`)
    @staticmethod
    def stepped(pg):
        return any(b.kind == "step" or (b.kind == "component" and b.data.get("build") is not None)
                   for b in pg.blocks)

    @staticmethod
    def split_layers(blocks):
        """Blocks → layers (`+++ overlay` starts one) → steps (`+++` starts one): [[{mark, blocks}]]."""
        layers = [[{"mark": {}, "blocks": []}]]
        for b in blocks:
            if b.kind == "step":
                if b.data["overlay"]:
                    layers.append([{"mark": b.data, "blocks": []}])
                else:
                    layers[-1].append({"mark": b.data, "blocks": []})
            else:
                layers[-1][-1]["blocks"].append(b)
        return layers

    def layer_frames(self, steps):
        """[(View, notes)] — one per slide of the layer. A `build` component expands its step into one slide
        per item; `+++ dim` / `+++ focus=` grey out everything from the earlier steps."""
        frames = []
        fade_before, focus = -1, None
        for k, st in enumerate(steps):
            m = st["mark"]
            if k and (m.get("dim") or m.get("focus")):
                fade_before, focus = k, m.get("focus")
            notes = m.get("notes") or []
            base = {"step": k, "fade_before": fade_before, "focus": focus}
            build = next((b for b in st["blocks"] if b.kind == "component" and b.data.get("build") is not None
                          and self.item_count(b.data)), None)
            if build is None:
                frames.append((View(**base), notes))
                continue
            n, mode = self.item_count(build.data), build.data["build"]
            if mode == "focus":  # overview first, then each item in turn
                frames.append((View(**base), notes))
                frames += [(View(**base, items={id(build): (n, j, None)}), notes) for j in range(n)]
            else:
                frames += [(View(**base, items={id(build): (j + 1, None, j if mode == "dim" else None)}), notes)
                           for j in range(n)]
        return frames

    def render_stepped(self, pg):
        """One slide per build step. Each layer is laid out once with all of its steps, so nothing moves
        between slides; an overlay layer is drawn over a veiled copy of the layers below."""
        layers = self.split_layers(pg.blocks)
        layers[0][0]["mark"] = {"notes": pg.notes}
        entries = [[(k, b) for k, st in enumerate(layer) for b in st["blocks"]] for layer in layers]
        frames = [self.layer_frames(layer) for layer in layers]
        sources = [b for layer in entries for _, b in layer if b.kind == "source"]
        top = self.content_top(pg)
        for L, layer_frames in enumerate(frames):
            for view, notes in layer_frames:
                s = self.new_slide(notes=notes or None)
                if pg.chapter is not None:
                    self.nav(s, pg.chapter, pg.sub)
                titles = None
                hole = None
                for j in range(L + 1):
                    if j == L:
                        v = view
                    else:  # a covered layer shows its finished state: every item revealed, no build focus
                        last = frames[j][-1][0]
                        v = View(step=last.step, fade_before=last.fade_before, focus=last.focus)
                    want = layers[j + 1][0]["mark"].get("focus") if j < L else None
                    lay, hole_ = self.draw_layer(s, pg, entries[j], v, bool(sources), overlay=j > 0, hole=hole,
                                                 want_box=want)
                    if titles is None:
                        titles = self.content_title(s, pg, lay)
                    if j < L:
                        hole = hole_
                        self.veil(s, (0, top - 0.1, W, H - top + 0.1), hole=hole)
                if L:
                    tree = s.shapes._spTree
                    for t in titles:  # keep the title crisp above the veil
                        tree.append(t._element)
                self.footer(s, sources)

    def footer(self, s, sources):
        if sources:
            src = "；".join(plain(b.data) for b in sources)
            self.text(s, PAD, H - 0.55, W - 2 * PAD - 1.2, 0.4, src, 9, color=self.muted,
                      anchor=MSO_ANCHOR.MIDDLE)
        self.counter_pill(s)

    def draw_layer(self, s, pg, entries, view, has_sources, overlay=False, hole=None, want_box=None):
        """Draw one layer of a stepped slide as `view` shows it. Returns (layout, box of the `want_box` items)."""
        text = [(k, b) for k, b in entries if b.kind in TEXT_KINDS]
        visuals = [(k, b) for k, b in entries if b.kind in ("image", "table", "code", "component")]
        callouts = [(k, b) for k, b in entries if b.kind == "callout"]
        paras = []
        for k, b in text:
            for pd in self.text_paras([b]):
                pd["step"] = k
                paras.append(pd)
        lay = self.layout(pg, paras, [b for _, b in visuals], [b for _, b in callouts], has_sources)
        if overlay and hole and callouts and not paras and not visuals:
            self.callouts_near(lay, pg, hole, has_sources)
        boxes = []
        x, y, w, h = lay["text"]
        statement = (not visuals and paras and len(paras) <= 4 and all(p["bullet"] == "none" for p in paras)
                     and any(f.get("hl") for p in paras for _, f in inline_runs(p["text"])))
        if lay["mode"] == "showcase":
            self.draw_showcase_text(s, pg, visuals[0][1].data, x, y, w, h)
        elif statement:
            tb = self.draw_statement(s, paras, x, y, w, h)
            for pd, p in zip(paras, tb.text_frame.paragraphs):
                if pd["step"] > view.step:
                    self.ghost(p._p)
                elif pd["step"] < view.fade_before:
                    self.fade(p._p)
        elif not visuals and len(text) == 1 and text[0][1].kind == "quote":
            if text[0][0] <= view.step:
                self.draw_quote(s, text[0][1].data, x, y, w, h)
        elif paras:
            self.draw_text(s, paras, x, y, w, h, pg.title, view=view, panel=overlay,
                           max_size=self.sizes["body_max"] - 4 if lay["mode"] == "below" else
                           (self.sizes["body_max"] + 4 if not visuals else None))
        if visuals:
            cells = self.visual_cells(len(visuals), *lay["visual"])
            for (k, v), cell in zip(visuals, cells):
                if k > view.step:
                    continue
                if overlay:
                    self.backdrop(s, v, cell)
                boxes.append(self.draw_visual(s, v, cell, pg.title, step=k, view=view, want_box=want_box))
        for (k, c), (cx, cy, cw, ch) in zip(callouts, lay["callouts"]):
            if k > view.step:
                continue
            start = len(s.shapes._spTree)
            self.callout(s, c.data, cx, cy, cw, ch)
            if k < view.fade_before:
                for el, _ in self.tagged(s, start):
                    self.fade(el)
        return lay, self.union(boxes)

    def callouts_near(self, lay, pg, hole, has_sources):
        """An overlay that only holds callouts annotates the spotlighted part: the callouts go right under
        the hole, or right above it when there is no room below."""
        _, hy, _, hh = hole
        boxes = lay["callouts"]
        total = boxes[-1][1] + boxes[-1][3] - boxes[0][1]
        top = self.content_top(pg) + 0.95
        bottom = H - (0.75 if has_sources else 0.68)
        for y0 in (hy + hh + 0.3, hy - 0.3 - total):
            if y0 >= top and y0 + total <= bottom:
                lay["callouts"] = [(cx, cy - boxes[0][1] + y0, cw, ch) for cx, cy, cw, ch in boxes]
                return

    def backdrop(self, s, v, cell):
        """Overlay visuals sit on a lifted panel so the veiled slide below does not show through them."""
        if v.kind in ("image", "table") or (v.kind == "component" and v.data["type"] in ("cards", "compare", "stats",
                                                                                        "chart")):
            return
        x, y, w, h = cell
        self.panel(s, x - 0.2, y - 0.15, w + 0.4, h + 0.3)

    def panel(self, s, x, y, w, h):
        sp = self.block(s, x, y, w, h, self.white, thin=True)
        if not self.elevation and self.surface in ("outline", "flat", "luxury", "ornate", "sketch"):
            self.soft_shadow(sp, blur=8, dist=3, angle=45, alpha=0.3)
        return sp

    def render_content(self, pg):
        s = self.new_slide(notes=pg.notes)
        text, visuals, callouts, sources = self.split_blocks(pg.blocks)
        if pg.chapter is not None:
            self.nav(s, pg.chapter, pg.sub)
        paras = self.text_paras(text)
        lay = self.layout(pg, paras, visuals, callouts, bool(sources))
        self.content_title(s, pg, lay)
        x, y, w, h = lay["text"]

        statement = (not visuals and paras and len(paras) <= 4 and all(p["bullet"] == "none" for p in paras)
                     and any(f.get("hl") for p in paras for _, f in inline_runs(p["text"])))
        only_quote = not visuals and len(text) == 1 and text[0].kind == "quote"
        if lay["mode"] == "showcase":
            self.draw_showcase_text(s, pg, visuals[0].data, x, y, w, h)
            self.draw_visuals(s, visuals, *lay["visual"], pg.title)
        elif statement:
            self.draw_statement(s, paras, x, y, w, h)
        elif only_quote:
            self.draw_quote(s, text[0].data, x, y, w, h)
        elif not visuals:
            self.draw_text(s, paras, x, y, w, h, pg.title, max_size=self.sizes["body_max"] + 4)
        elif not paras:
            self.draw_visuals(s, visuals, *lay["visual"], pg.title)
        else:
            self.draw_text(s, paras, x, y, w, h, pg.title,
                           max_size=self.sizes["body_max"] - 4 if lay["mode"] == "below" else None)
            self.draw_visuals(s, visuals, *lay["visual"], pg.title)

        for c, (cx, cy, cw, ch) in zip(callouts, lay["callouts"]):
            self.callout(s, c.data, cx, cy, cw, ch)
        self.footer(s, sources)

    def draw_text(self, s, paras, x, y, w, h, title, max_size=None, view=None, panel=False):
        """Body text, sized to fit. With a build `view`, the size is chosen for all paragraphs but only the
        steps shown so far are written (text is top-anchored, so earlier lines never move)."""
        if not paras:
            return
        size = self.sizes["body_min"]
        for sz in range(int(max_size or self.sizes["body_max"]), int(self.sizes["body_min"]) - 1, -1):
            if estimate_height_pt(paras, sz * 1.08, E(w)) <= h * 72:
                size = sz
                break
        else:
            self.warnings.append(f"[{title}] text may overflow — split the slide or trim it")
        if panel:
            th = min(h, estimate_height_pt(paras, size * 1.08, E(w)) / 72)
            self.panel(s, x - 0.25, y - 0.2, w + 0.5, th + 0.4)
        tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
        tf = tb.text_frame
        self.tf_setup(tf)
        num = {}
        if view is not None:
            paras = [pd for pd in paras if pd.get("step", 0) <= view.step]
        for k, pd in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            lvl = pd.get("level", 0)
            sz = max(10, round(size * pd.get("scale", 1.0)))
            p.space_after = Pt(round(sz * 0.55))
            p.line_spacing = 1.12
            pPr = cjk_breaks(p)
            if pd["bullet"] == "sq":
                ind = 0.32 + 0.3 * lvl
                pPr.set("marL", str(E(ind)))
                pPr.set("indent", str(-E(0.28)))
                if pd.get("ordered"):
                    num[lvl] = num.get(lvl, 0) + 1
                    etree.SubElement(pPr, qn("a:buFont"), typeface=self.f_display)
                    etree.SubElement(pPr, qn("a:buAutoNum"), type="arabicPeriod", startAt=str(num[lvl]))
                else:
                    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
                    etree.SubElement(pPr, qn("a:buChar"), char="■" if lvl == 0 else "–")
            else:
                num.clear()
                pPr.set("marL", "0")
                pPr.set("indent", "0")
                etree.SubElement(pPr, qn("a:buNone"))
            self.runs(p, pd["text"], sz, bold=pd.get("bold"), color=pd.get("color", self.ink),
                      italic=pd.get("italic"))
            if view is not None and pd.get("step", 0) < view.fade_before:
                self.fade(p._p)
        return tb

    def draw_showcase_text(self, s, pg, d, x, y, w, h):
        """Left column of an image showcase: the subtitle becomes the headline, the caption sits under it."""
        lines = []
        if pg.subtitle:
            avail = (w - 0.3) * 72
            size = min(34, max(20, int(avail / max(1.0, disp_width(pg.subtitle)) * 0.95)))
            clauses = [c for c in re.split(r"(?<=[，、；：。,;:])\s*", pg.subtitle) if c]
            if len(clauses) > 1 and disp_width(pg.subtitle) * size > avail:
                # break at the punctuation instead of mid-phrase; size to the longest clause
                size = min(34, max(20, int(avail / max(disp_width(c) for c in clauses) * 0.95)))
            else:
                clauses = [pg.subtitle]
            for k, c in enumerate(clauses):
                lines.append({"text": c, "size": size, "bold": True, "color": self.title_ink, "font": "heading",
                              "space_after": 14 if k == len(clauses) - 1 else 2})
        if d.get("alt"):
            lines.append({"text": d["alt"], "size": 16, "color": self.muted})
        self.text(s, x, y, w - 0.3, h, lines, 30, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)

    def draw_statement(self, s, paras, x, y, w, h):
        lines = [{"text": p["text"], "size": 28, "bold": True, "space_after": 22} for p in paras]
        tb = self.text(s, x + 0.6, y, w - 1.2, h, lines, 28, color=self.title_ink, anchor=MSO_ANCHOR.MIDDLE,
                       align=PP_ALIGN.CENTER, font="heading")
        self.decorate(s, (x + w - 1.3, y, 1.3, 1.3), "accent")
        return tb

    def draw_quote(self, s, text, x, y, w, h):
        cw, ch = min(w, 10.0), min(h, 3.2)
        cx, cy = x + (w - cw) / 2, y + (h - ch) / 2
        self.block(s, cx, cy, cw, ch, self.white)
        self.text(s, cx + 0.5, cy + 0.3, cw - 1.0, ch - 0.6, f"「{text}」", 30, bold=True, color=self.card_ink,
                  anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, font="heading")
        self.pill(s, cx - 0.2, cy - 0.3, "QUOTE", self.pal(0), size=12, rot=-5)

    CALLOUT = {"note": ("註", 1), "tip": ("TIP", 2), "important": ("重點", 0), "warning": ("注意", 3),
               "caution": ("注意", 0), "summary": ("小結", 2), "problem": ("問題", 0), "result": ("結果", 2),
               "success": ("結果", 2), "hint": ("HINT", 3)}

    def callout_pill(self, s, d, x, y, w, h):
        label, ci = self.CALLOUT.get(d["kind"], (d["kind"].upper(), 1))
        fill = self.tone_fill(d["kind"]) or self.color_for(ci)
        self.block(s, x, y, w, h, fill, thin=True)
        size = 15
        while size > 11 and self.lines_height([{"text": d["text"], "size": size}], w - 1.6) + 0.2 > h:
            size -= 1
        self.text(s, x + 1.3, y, w - 1.6, h, d["text"], size, bold=True, color=self.on(fill),
                  anchor=MSO_ANCHOR.MIDDLE)
        self.pill(s, x + 0.2, y + (h - 0.44) / 2, label, self.white, size=13, rot=-4)

    # ------------------------------------------------------------------ visuals
    @staticmethod
    def visual_cells(n, x, y, w, h):
        gap = 0.4
        if n <= 1:
            return [(x, y, w, h)]
        cw = (w - gap * (n - 1)) / n
        return [(x + k * (cw + gap), y, cw, h) for k in range(n)]

    def draw_visuals(self, s, visuals, x, y, w, h, title):
        for v, cell in zip(visuals, self.visual_cells(len(visuals), x, y, w, h)):
            self.draw_visual(s, v, cell, title)

    def draw_visual(self, s, v, cell, title, step=0, view=None, want_box=None):
        """Draw one visual, then apply its build state: hide items not revealed yet, grey out what is not in
        focus (`focus=` on the component, `+++ focus=` / `+++ dim` steps). Returns the box of the items named
        by `want_box` (the spotlight of an overlay above), or None."""
        d = v.data
        faded = view is not None and step < view.fade_before
        st = view.items.get(id(v)) if view is not None else None
        if v.kind == "component" and d["type"] == "diagram" and faded and view.focus:
            info = re.sub(r"(?:^|\s)focus=.+?(?=\s+\w+=|$)", "", d.get("info", ""))  # the step's focus wins
            d = dict(d, info=f"{info} focus={view.focus}")  # diagrams draw their own focus
            faded = False
        start = len(s.shapes._spTree)
        if v.kind == "component":
            getattr(self, f"comp_{d['type'].replace('timeline', 'steps')}")(s, d, *cell)
        else:
            getattr(self, f"draw_{v.kind}")(s, d, *cell, title=title)
        els = self.tagged(s, start)
        self._marks = [m for m in self._marks if m[1] <= start]

        if st:  # build: drop the items not revealed yet
            for el, tg in els:
                if tg.get("item") is not None and tg["item"] >= st[0]:
                    el.getparent().remove(el)
            els = [(el, tg) for el, tg in els if el.getparent() is not None]
        if faded and view.focus:
            spec = view.focus
        elif st and st[1] is not None:
            spec = str(st[1] + 1)
        elif v.kind == "component" and d["type"] != "diagram":
            spec = d.get("focus")
        else:
            spec = None
        dim_all = faded and not view.focus
        idx, names = parse_focus(spec)
        rows = self.table_rows(v, spec) if v.kind == "table" else set()
        charts = []
        for el, tg in els:
            if tg.get("region") and ((spec and not in_focus(tg, idx, names)) or
                                     (tg.get("focus_only") and not spec)):
                el.getparent().remove(el)
                continue
            if tg.get("fixed") and not dim_all:
                continue
            if v.kind == "table" and spec and not dim_all:
                tbl = el.find(".//" + qn("a:tbl"))
                if tbl is not None:
                    for r, tr in enumerate(tbl.findall(qn("a:tr"))):
                        if r and r not in rows:
                            self.fade(tr)
                    continue
            if dim_all or (spec and not in_focus(tg, idx, names)) or \
                    (st and st[2] is not None and tg.get("item") is not None and tg["item"] < st[2]):
                if self.fade(el) == "chart":
                    charts.append(self.el_box(el))
        for box in charts:
            self.veil(s, box)
        if not want_box:
            return None
        if v.kind == "table":  # exact row band: padding would cut into the neighbouring rows
            return self.row_box(els, self.table_rows(v, want_box))
        w_idx, w_names = parse_focus(want_box)
        box = self.union(self.el_box(el) for el, tg in els if el.getparent() is not None and
                         in_focus(tg, w_idx, w_names))
        return (box[0] - 0.12, box[1] - 0.12, box[2] + 0.24, box[3] + 0.24) if box else None

    @staticmethod
    def table_rows(v, spec):
        """Body rows (1-based, the header is row 0) named by a focus spec: numbers or first-cell text."""
        idx, names = parse_focus(spec)
        rows = v.data["rows"]
        return {i + 1 for i in idx} | {r for r in range(1, len(rows)) if plain(rows[r][0]).strip() in names}

    def row_box(self, els, rows):
        for el, _ in els:
            tbl = el.find(".//" + qn("a:tbl"))
            if tbl is None:
                continue
            x, y, w, _ = self.el_box(el)
            boxes, yy = [], y
            for r, tr in enumerate(tbl.findall(qn("a:tr"))):
                rh = int(tr.get("h")) / IN
                if r in rows:
                    boxes.append((x, yy, w, rh))
                yy += rh
            return self.union(boxes)
        return None

    def draw_image(self, s, d, x, y, w, h, title=""):
        from PIL import Image
        p = Path(d["path"])
        p = p if p.is_absolute() else self.base / p
        if not p.exists():
            self.warnings.append(f"[{title}] image not found: {d['path']}")
            return
        svg = p.suffix.lower() == ".svg"
        cap = 0.5 if d.get("alt") else 0.0
        if svg:
            data = p.read_bytes()
            iw, ih = svg_aspect(data), 1.0
        else:
            with Image.open(p) as im:
                iw, ih = im.size
        sc = min((w - 0.05) / iw, (h - cap - 0.05) / ih)
        pw, ph = iw * sc, ih * sc
        px, py = x + (w - pw) / 2, y + (h - cap - ph) / 2
        if not svg and iw / max(0.1, pw) < 90:
            self.warnings.append(f"[{title}] image is low resolution: {iw}px across {pw:.1f}in "
                                 f"(~{iw / max(0.1, pw):.0f} dpi) — use a larger source")
        box = (E(px), E(py), E(pw), E(ph))
        if svg:
            pic = add_svg_picture(s, data, *box, descr=d.get("alt") or p.stem)
        else:
            pic = s.shapes.add_picture(str(p), *box)
        self.outline(pic)
        if d.get("alt"):
            self.pill(s, px, py + ph + 0.1, d["alt"], self.pal(3), size=11, upper=False,
                      font="body")

    def draw_code(self, s, d, x, y, w, h, title=""):
        lines = d["text"].split("\n")
        size = 16
        longest = max((len(l) for l in lines), default=1)
        while size > 9 and (longest * size * 0.6 / 72 > w - 0.6 or len(lines) * size * 1.3 / 72 > h - 0.6):
            size -= 1
        ch = min(h, len(lines) * size * 1.3 / 72 + 0.7)
        cy = y + (h - ch) / 2
        self.block(s, x, cy, w, ch, "1A1A1A")
        self.text(s, x + 0.3, cy + 0.35, w - 0.6, ch - 0.5,
                       [{"text": l or " ", "font": "code"} for l in lines], size, font="code", color="F5F5F5")
        self.pill(s, x + 0.2, cy - 0.2, d.get("lang") or "CODE", self.pal(3), size=11, rot=-3)

    def draw_table(self, s, d, x, y, w, h, title="", header_color=None):
        rows = d["rows"]
        nc = max(len(r) for r in rows)
        rows = [r + [""] * (nc - len(r)) for r in rows]
        nr = len(rows)
        size = max(10, min(16, int(22 - 0.8 * nr - 0.6 * nc)))
        rh = size * 2.2 / 72
        th = min(h, rh * nr)
        tx = x
        gf = s.shapes.add_table(nr, nc, E(tx), E(y), E(w), E(th))
        # shadow must sit behind the table
        if self.surface in ("brutal", "pixel"):
            shadow = self.shape(s, MSO_SHAPE.RECTANGLE, tx + self.sw / 72, y + self.sw / 72, w, th, fill=self.black)
            gf._element.addprevious(shadow._element)
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        tblPr.set("bandRow", "0")
        tblPr.set("firstRow", "1")
        widths = [max(disp_width(plain(r[c])) for r in rows) + 2 for c in range(nc)]
        tot = sum(widths)
        for c in range(nc):
            tbl.columns[c].width = E(w * widths[c] / tot)
        aligns = d.get("aligns") or []
        amap = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
        hc = header_color or self.pal(2)
        for r, row in enumerate(rows):
            tbl.rows[r].height = E(rh)
            for c, val in enumerate(row):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cf = hc if r == 0 else (self.white if r % 2 else self.ground)
                cell.fill.fore_color.rgb = rgb(cf)
                cell.margin_left = cell.margin_right = E(0.08)
                cell.margin_top = cell.margin_bottom = E(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tcPr = cell._tc.get_or_add_tcPr()
                for k, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                    ln = etree.Element(qn(tag), w=str(int(self.table_rule * 12700)), cap="flat", cmpd="sng", algn="ctr")
                    sf = etree.SubElement(ln, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr"), val=self.table_line)
                    tcPr.insert(k, ln)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                v = val.strip()
                if c < len(aligns) and aligns[c] != "l":
                    p.alignment = amap[aligns[c]]
                elif r == 0 or re.fullmatch(r"[○●◎×✓✗✔✘△\-—vVxXoO]|[\d.,%+\-]+", v or "-"):
                    p.alignment = PP_ALIGN.CENTER
                self.runs(p, val, size, bold=True if r == 0 else None, color=self.on(cf))
        if rh * nr > h * 1.05:
            self.warnings.append(f"[{title}] table may overflow ({nr} rows)")

    # ------------------------------------------------------------------ components
    def item_icon(self, s, x, y, d, k, it, size=18):
        fill = self.color_for(k + 1)
        self.block(s, x, y, d, d, fill, thin=True)
        label = it["attrs"].get("icon") or f"{k + 1}"
        if self.glyph(s, label, x, y, d, d, self.on(fill)):
            return
        font = "display" if label.isascii() else "body"
        self.text(s, x, y, d, d, label, size, font=font, color=self.on(fill), align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)

    def comp_cards(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        args = d["args"]
        grid = "grid" in args or ("stack" not in args and w > 8)
        gap = 0.35
        if not grid:
            ch = min(1.6, (h - gap * (n - 1)) / n)
            y0 = y + (h - (ch * n + gap * (n - 1))) / 2
            for k, it in enumerate(items):
                with self.tag(s, item=k, name=plain(it["title"])):
                    self.card(s, it, k, x + 0.25, y0 + k * (ch + gap), w - 0.35, ch, horizontal=True)
        else:
            cols = int(next((a.split("=")[1] for a in args if a.startswith("cols=")), 0)) or (
                n if n <= 3 else 2 if n == 4 else 3)
            rows = -(-n // cols)
            cw = (w - gap * (cols - 1)) / cols
            need = max(self.lines_height(self.card_lines(it), cw - 0.6) for it in items) + 0.62 + 0.9
            ch = min(max(need, 2.6), (h - gap * (rows - 1)) / rows)
            y0 = y + (h - (ch * rows + gap * (rows - 1))) / 2
            for k, it in enumerate(items):
                r, c = divmod(k, cols)
                with self.tag(s, item=k, name=plain(it["title"])):
                    self.card(s, it, k, x + c * (cw + gap), y0 + r * (ch + gap), cw - 0.1, ch, horizontal=False)

    # ------------------------------------------------------------------ tinted boxes (`style.card: tint`)
    def accent_ink(self, base):
        """A readable version of an accent colour for icons / numbers on its own light tint."""
        return base if self.luminance(base) < 0.5 else mix(base, "000000", 0.35)

    def tint_box(self, s, x, y, w, h, base=None, lw=2.0, fill=None, kind=None):
        """Rounded box in a light tint of `base` with a `base` outline (lifted when the theme has elevation)."""
        base = base or self.card_accent
        kind = kind or (MSO_SHAPE.ROUNDED_RECTANGLE if self.radius > 0 else MSO_SHAPE.RECTANGLE)
        sp = self.shape(s, kind, x, y, w, h, fill=fill or mix(base, self.ground, 0.88), line=base, lw=lw)
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = min(0.5, max(self.radius, 0.08) * 1.6 / max(0.05, min(w, h)))
        return self.lift(sp)

    def badge_at(self, s, cx, cy, d, label, base=None, size=None):
        """Circular number / icon badge centred on (cx, cy) — sits on a card's corner with `badge: corner`."""
        base = base or self.card_accent
        sp = self.shape(s, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=mix(base, self.ground, 0.78),
                        line=base, lw=1.75)
        self.lift(sp, small=True)
        ink = self.accent_ink(base)
        if not self.glyph(s, label, cx - d / 2, cy - d / 2, d, d, ink, scale=0.56):
            self.text(s, 0, 0, 0, 0, label, size or max(10, int(d * 30)), font="display", bold=True,
                      color=self.card_ink, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, target=sp)
        return sp

    def legend_tab(self, s, x, y, text, color, icon=None, size=12, fill=None):
        """Label that sits on a box's top border like a fieldset legend (centre line at y)."""
        tw = self.measure(text, size) + (0.42 if icon else 0.2)
        h = size / 72 * 1.9
        self.shape(s, MSO_SHAPE.RECTANGLE, x, y - h / 2, tw, h, fill=fill or self.ground)
        tx = x + 0.08
        if icon:
            self.glyph(s, icon, tx, y - h / 2, h, h, color, scale=0.72)
            tx += h * 0.8
        self.text(s, tx, y - h / 2, tw, h, text, size, bold=True, color=color, anchor=MSO_ANCHOR.MIDDLE)
        return tw

    def card_tint(self, s, it, k, x, y, w, h, horizontal):
        base = self.card_accent
        self.tint_box(s, x, y, w, h, base)
        icon = it["attrs"].get("icon")
        corner = self.badge == "corner"
        d = 0.46
        lines = self.card_lines(it)
        if corner:
            self.badge_at(s, x + 0.1, y + 0.1, d, f"{k + 1}", base)
        if horizontal:
            ix = x + (0.35 if corner else 0.25)
            if not corner:
                self.badge_at(s, ix + d / 2, y + h / 2, d, f"{k + 1}", base)
                ix += d + 0.2
            if icon and self.glyph(s, icon, ix, y + 0.15, min(0.9, h - 0.3), h - 0.3, self.accent_ink(base), 0.8):
                ix += min(0.9, h - 0.3) + 0.15
            self.text(s, ix, y + 0.1, x + w - ix - 0.2, h - 0.2, lines, 20, anchor=MSO_ANCHOR.MIDDLE)
        else:
            top = y + (0.42 if corner else 0.25)
            if not corner:
                self.badge_at(s, x + 0.25 + d / 2, top + d / 2, d, f"{k + 1}", base)
                top += d + 0.15
            align = PP_ALIGN.LEFT
            if icon:
                ih = min(1.1, h * 0.32)
                if self.glyph(s, icon, x, top, w, ih, self.accent_ink(base), 0.9):
                    top += ih + 0.15
                    align = PP_ALIGN.CENTER
            self.text(s, x + 0.25, top, w - 0.5, y + h - top - 0.15, lines, 20, align=align)
        if it["attrs"].get("tag"):
            if self.legend:
                self.legend_tab(s, x + 0.55, y, it["attrs"]["tag"], self.accent_ink(base),
                                fill=self.ground)
            else:
                self.pill(s, x + w - 1.2, y - 0.2, it["attrs"]["tag"], base, size=11, upper=False, font="body",
                          shadow=False)

    def card(self, s, it, k, x, y, w, h, horizontal):
        if self.tinted:
            return self.card_tint(s, it, k, x, y, w, h, horizontal)
        self.block(s, x, y, w, h, self.white)
        img = it["attrs"].get("img")
        inner_x, inner_w = x + 0.25, w - 0.5
        if img:
            iw = min(h - 0.3, w * 0.3)
            self.picture(s, img, x + w - iw - 0.18, y + (h - iw) / 2, iw, iw)
            inner_w -= iw + 0.15
        lines = self.card_lines(it)
        if horizontal:
            d = min(0.62, h - 0.4)
            self.item_icon(s, inner_x, y + (h - d) / 2, d, k, it)
            self.text(s, inner_x + d + 0.25, y + 0.1, inner_w - d - 0.25, h - 0.2, lines, 20,
                      anchor=MSO_ANCHOR.MIDDLE)
        else:
            d = 0.62
            self.item_icon(s, inner_x, y + 0.3, d, k, it)
            self.text(s, inner_x, y + 0.3 + d + 0.25, inner_w, h - d - 0.7, lines, 20)
        if it["attrs"].get("tag"):
            self.pill(s, x - 0.3, y - 0.28, it["attrs"]["tag"], self.color_for(k + 3), size=12, rot=-8,
                      upper=False, font="body")

    def comp_steps(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        horizontal = d["type"] == "timeline" or "horizontal" in d["args"] or (w > 8 and n <= 5
                                                                                and "vertical" not in d["args"])
        if "axis" in d["args"]:
            return self.steps_axis(s, items, x, y, w, h)
        if any(it.get("panels") for it in items):
            if horizontal:
                return self.steps_panels(s, d, x, y, w, h)
            items = [dict(it, children=it["children"] + [line for p in it.get("panels", [])
                                                       for line in [p["title"], *p["lines"]]]) for it in items]
        if self.tinted:
            return self.steps_tint(s, items, x, y, w, h, horizontal)
        if horizontal:
            gap = 0.55
            cw = (w - gap * (n - 1)) / n
            need = 0.55 + 0.4 + max(self.lines_height(
                [{"text": it["title"], "size": 20, "space_after": 6}, {"text": it["desc"] or "", "size": 15}]
                + [{"text": "■ " + c, "size": 14} for c in it["children"]], cw - 0.4) for it in items)
            ch = min(h, max(2.8, need))
            cy = y + (h - ch) / 2
            for k in range(n - 1):
                x1 = x + (k + 1) * cw + k * gap
                with self.tag(s, item=k + 1):
                    self.line(s, x1, cy + ch / 2, x1 + gap, cy + ch / 2, width=4)
            for k, it in enumerate(items):
                with self.tag(s, item=k, name=plain(it["title"])):
                    self.step_card(s, it, k, x + k * (cw + gap), cy, cw, ch)
            return
        step = min(1.35, h / n)
        y0 = y + (h - step * n) / 2
        dsz = min(0.6, step * 0.62)
        lx = x + 0.2 + dsz / 2
        if n > 1:
            self.line(s, lx, y0 + step / 2, lx, y0 + step * (n - 0.5), width=3)
        for k, it in enumerate(items):
            with self.tag(s, item=k, name=plain(it["title"])):
                self.step_row(s, it, k, x, y0 + k * step + step / 2, w, step, lx, dsz)

    def steps_panels(self, s, d, x, y, w, h):
        items = d["items"]
        top_h, gap = h * 0.5, 0.3
        start = len(s.shapes._spTree)
        clean = [dict(it, panels=[], extra=[], children=it["children"] + it["extra"]) for it in items]
        self.comp_steps(s, dict(d, items=clean), x, y, w, top_h)
        els = self.tagged(s, start)
        boxes = {k: self.union(self.el_box(el) for el, tg in els if tg.get("item") == k) for k in range(len(items))}
        panels = [(k, p) for k, it in enumerate(items) for p in it.get("panels", [])]
        pw = (w - gap * (len(panels) - 1)) / len(panels)
        py, ph = y + top_h + gap, h - top_h - gap
        for j, (k, panel) in enumerate(panels):
            px = x + j * (pw + gap)
            with self.tag(s, item=k, name=plain(items[k]["title"])):
                box = boxes[k]
                if box:
                    self.line(s, box[0] + box[2] / 2, box[1] + box[3], px + pw / 2, py,
                              color=self.muted, width=0.75)
                fill = mix(self.card_accent, self.ground, 0.9)
                self.shape(s, MSO_SHAPE.RECTANGLE, px, py, pw, ph, fill=fill)
                lines = [{"text": panel["title"], "size": 16, "bold": True, "color": self.on(fill), "space_after": 6}]
                lines += [{"text": str(i + 1) + ". " + re.sub(r"^\d+[.)]\s*", "", line), "size": 13,
                           "color": self.on(fill)} for i, line in enumerate(panel["lines"])]
                while self.lines_height(lines, pw - 0.4) > ph - 0.3 and lines[0]["size"] > 10:
                    for line in lines:
                        line["size"] -= 1
                self.text(s, px + 0.2, py + 0.15, pw - 0.4, ph - 0.3, lines, 16)

    def steps_axis(self, s, items, x, y, w, h):
        n = len(items)
        cw = w / n
        ay = y + 0.6
        photos = any(it["attrs"].get("img") for it in items)
        lines_by_item = []
        for it in items:
            lines = [{"text": it["desc"], "size": 16, "bold": True, "color": self.title_ink, "space_after": 5}]
            lines += [{"text": t, "size": 12, "color": self.muted} for t in it["extra"] + it["children"] if t]
            lines_by_item.append(lines)
        th = max(self.lines_height(lines, cw - 0.24) for lines in lines_by_item) + 0.16
        photo_y = ay + 0.25 + th + 0.15
        for k, (it, lines) in enumerate(zip(items, lines_by_item)):
            cx = x + (k + 0.5) * cw
            fill = self.card_accent if self.tinted else self.color_for(k)
            with self.tag(s, item=k, name=plain(it["title"])):
                self.line(s, x + k * cw, ay, x + (k + 1) * cw, ay,
                          color=mix(self.card_accent, self.ground, 0.55), width=1.5)
                self.shape(s, MSO_SHAPE.OVAL, cx - 0.075, ay - 0.075, 0.15, 0.15, fill=fill)
                size = 18
                while size > 10 and self.measure(it["title"], size, "display") > cw - 0.2:
                    size -= 1
                self.text(s, x + k * cw + 0.1, y, cw - 0.2, 0.48, it["title"], size, bold=True,
                          color=self.accent_ink(fill), font="display", align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE)
                self.text(s, x + k * cw + 0.12, ay + 0.25, cw - 0.24, th, lines, 16, align=PP_ALIGN.CENTER)
                if photos and it["attrs"].get("img") and y + h > photo_y:
                    pic = self.picture(s, it["attrs"]["img"], x + k * cw + 0.12, photo_y,
                                       cw - 0.24, y + h - photo_y)
                    if pic:
                        self.outline(pic[0])

    def steps_tint(self, s, items, x, y, w, h, horizontal):
        """Tinted cards with corner number badges joined by accent arrows (or a badge column when vertical)."""
        n = len(items)
        base = self.card_accent
        if horizontal:
            gap = 0.75
            cw = (w - gap * (n - 1)) / n
            ch = min(h, max(2.4, 0.9 + max(self.lines_height(self.card_lines(it, 19, 14, 13), cw - 0.5)
                                           + (1.0 if it["attrs"].get("icon") else 0) for it in items)))
            cy = y + (h - ch) / 2
            for k in range(n - 1):
                x1 = x + (k + 1) * cw + k * gap
                with self.tag(s, item=k + 1):
                    self.shape(s, MSO_SHAPE.RIGHT_ARROW, x1 + 0.16, cy + ch / 2 - 0.13, gap - 0.32, 0.26, fill=base)
            for k, it in enumerate(items):
                cx = x + k * (cw + gap)
                with self.tag(s, item=k, name=plain(it["title"])):
                    self.tint_box(s, cx, cy, cw, ch, base)
                    self.badge_at(s, cx + 0.12, cy + 0.12, 0.5, f"{k + 1}", base)
                    top = cy + 0.45
                    icon = it["attrs"].get("icon")
                    if icon and self.glyph(s, icon, cx, top, cw, 0.85, self.accent_ink(base), 0.95):
                        top += 1.0
                    self.text(s, cx + 0.25, top, cw - 0.5, cy + ch - top - 0.15, self.card_lines(it, 19, 14, 13), 19,
                              align=PP_ALIGN.CENTER if icon else PP_ALIGN.LEFT)
            return
        step = min(1.3, h / n)
        y0 = y + (h - step * n) / 2
        d = min(0.62, step * 0.6)
        for k, it in enumerate(items):
            cy = y0 + k * step + step / 2
            with self.tag(s, item=k, name=plain(it["title"])):
                self.badge_at(s, x + 0.2 + d / 2, cy, d, it["attrs"].get("icon") or f"{k + 1:02d}", base,
                              size=int(d * 26))
                tx, tw = x + d + 0.55, w - d - 0.75
                head = {"text": it["title"], "size": 19, "bold": True, "color": self.title_ink}
                body = [{"text": it["desc"], "size": 15, "color": self.muted}] if it["desc"] else []
                body += [{"text": "■ " + c, "size": 13, "color": self.muted} for c in it["children"]]
                hh = self.lines_height([head], tw)
                bh = self.lines_height(body, tw) if body else 0
                ty = cy - (hh + 0.14 + bh) / 2
                self.text(s, tx, ty, tw, hh + 0.05, [head], 19)
                self.shape(s, MSO_SHAPE.RECTANGLE, tx, ty + hh + 0.03, min(tw, self.measure(it["title"], 19) + 0.2),
                           0.045, fill=base)
                if body:
                    self.text(s, tx, ty + hh + 0.14, tw, bh + 0.1, body, 15)

    def step_card(self, s, it, k, cx, cy, cw, ch):
        self.block(s, cx, cy, cw, ch, self.white, thin=True)
        self.shape(s, MSO_SHAPE.RECTANGLE, cx, cy, cw, 0.55, fill=self.color_for(k), line=self.black, lw=self.tw)
        has_icon = self.glyph(s, it["attrs"].get("icon"), cx + 0.1, cy + 0.05, 0.45, 0.45,
                              self.on(self.color_for(k)), scale=0.75)
        self.text(s, cx + (0.4 if has_icon else 0), cy, cw - (0.4 if has_icon else 0), 0.55,
                  f"STEP {k + 1}", 15, font="label", bold=True, color=self.on(self.color_for(k)),
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spc=1.5)
        lines = [{"text": it["title"], "size": 20, "bold": True, "color": self.card_ink, "space_after": 6}]
        if it["desc"]:
            lines.append({"text": it["desc"], "size": 15, "color": self.card_muted})
        lines += [{"text": "■ " + c, "size": 14, "color": self.card_muted} for c in it["children"]]
        self.text(s, cx + 0.2, cy + 0.75, cw - 0.4, ch - 0.9, lines, 20)

    def step_row(self, s, it, k, x, cy, w, step, lx, dsz):
        self.block(s, lx - dsz / 2, cy - dsz / 2, dsz, dsz, self.color_for(k), thin=True)
        if not self.glyph(s, it["attrs"].get("icon"), lx - dsz / 2, cy - dsz / 2, dsz, dsz,
                          self.on(self.color_for(k))):
            self.text(s, lx - dsz / 2, cy - dsz / 2, dsz, dsz, str(k + 1), 18, font="display",
                      color=self.on(self.color_for(k)), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        lines = [{"text": it["title"], "size": 20, "bold": True, "color": self.title_ink, "space_after": 2}]
        if it["desc"]:
            lines.append({"text": it["desc"], "size": 15, "color": self.muted})
        lines += [{"text": "■ " + c, "size": 14, "color": self.muted} for c in it["children"]]
        self.text(s, lx + dsz / 2 + 0.3, cy - step / 2, w - dsz - 0.7, step, lines, 20, anchor=MSO_ANCHOR.MIDDLE)

    def comp_flow(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        vertical = "vertical" in d["args"] or (w / h < 1.3 and "horizontal" not in d["args"])
        aw = 0.42
        if vertical:
            bh = min(1.25, (h - aw * (n - 1)) / n)
            bw = min(w, 4.6)
            bx = x + (w - bw) / 2
            y0 = y + (h - (bh * n + aw * (n - 1))) / 2
            for k, it in enumerate(items):
                by = y0 + k * (bh + aw)
                with self.tag(s, item=k, name=plain(it["title"])):
                    self.flow_box(s, it, k, bx, by, bw, bh)
                if k < n - 1:
                    with self.tag(s, item=k + 1):
                        self.shape(s, MSO_SHAPE.DOWN_ARROW, bx + bw / 2 - 0.17, by + bh + 0.06, 0.34, aw - 0.12,
                                   fill=self.arrow_fill())
            return
        per_row = n if n <= 6 else -(-n // 2)
        rows = -(-n // per_row)
        bw = (w - aw * (per_row - 1)) / per_row
        bh = min(2.2, (h - 0.5 * (rows - 1)) / rows)
        y0 = y + (h - (bh * rows + 0.5 * (rows - 1))) / 2
        for k, it in enumerate(items):
            r, c = divmod(k, per_row)
            bx, by = x + c * (bw + aw), y0 + r * (bh + 0.5)
            with self.tag(s, item=k, name=plain(it["title"])):
                self.flow_box(s, it, k, bx, by, bw - 0.05, bh)
            if c < per_row - 1 and k < n - 1:
                with self.tag(s, item=k + 1):
                    self.shape(s, MSO_SHAPE.RIGHT_ARROW, bx + bw + 0.02, by + bh / 2 - 0.16, aw - 0.12, 0.32,
                               fill=self.arrow_fill())

    def arrow_fill(self):
        return self.card_accent if self.tinted else self.black

    def flow_box(self, s, it, k, x, y, w, h):
        if self.tinted:
            self.tint_box(s, x, y, w, h, self.card_accent)
            fill, ink, icon_ink = mix(self.card_accent, self.ground, 0.88), self.card_ink, self.accent_ink(self.card_accent)
        else:
            fill = self.color_for(k)
            self.block(s, x, y, w, h, fill, thin=True)
            ink = icon_ink = self.on(fill)
        top = 0.0
        if h >= 1.3 and self.glyph(s, it["attrs"].get("icon"), x, y + 0.12, w, 0.55, icon_ink, scale=0.95):
            top = 0.55
        lines = [{"text": it["title"], "size": 19, "bold": True, "color": ink, "space_after": 3}]
        if it["desc"]:
            lines.append({"text": it["desc"], "size": 14, "color": self.card_muted if self.tinted else ink})
        self.text(s, x + 0.12, y + 0.05 + top, w - 0.24, h - 0.1 - top, lines, 19, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)

    def comp_stats(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        cols = n if (w > 8 or n <= 2) else 2
        rows = -(-n // cols)
        gap = 0.45
        cw = (w - gap * (cols - 1)) / cols
        ch = min(3.0, (h - gap * (rows - 1)) / rows)
        y0 = y + (h - (ch * rows + gap * (rows - 1))) / 2
        for k, it in enumerate(items):
            r, c = divmod(k, cols)
            with self.tag(s, item=k, name=plain(it["title"])):
                self.stat_card(s, it, k, x + c * (cw + gap), y0 + r * (ch + gap), cw, ch)

    def stat_card(self, s, it, k, cx, cy, cw, ch):
        rot = (-2 if k % 2 == 0 else 2) if self.tilt else 0
        if self.tinted:
            self.tint_box(s, cx, cy, cw, ch, self.card_accent)
        else:
            self.block(s, cx, cy, cw, ch, self.white, thin=True, rot=rot)
        if self.deco:
            self.shape(s, MSO_SHAPE.OVAL, cx + cw - 0.42, cy + 0.2, 0.2, 0.2, fill=self.color_for(k),
                       line=self.black, lw=self.tw)
        if it["attrs"].get("img"):
            paths = [p.strip() for p in it["attrs"]["img"].split(",") if p.strip()]
            iw = (cw - 0.6) / len(paths)
            for i, path in enumerate(paths):
                self.picture(s, path, cx + 0.3 + i * iw, cy + 0.25, iw - 0.08, ch * 0.48)
            lines = [{"text": it["title"], "size": 20, "bold": True, "color": self.card_ink}]
            lines += [{"text": t, "size": 14, "color": self.card_muted}
                      for t in [it["desc"], *it["extra"], *it["children"]] if t]
            self.text(s, cx + 0.3, cy + ch * 0.58, cw - 0.6, ch * 0.42 - 0.15, lines, 20)
            return
        num = it["title"]
        size = 60 if len(num) <= 5 else 44 if len(num) <= 8 else 32
        ink = self.accent_ink(self.card_accent) if self.tinted else self.stat_ink(k)
        lines = [{"text": num, "size": size, "font": "display", "color": ink, "space_after": 4}]
        if it["desc"]:
            lines.append({"text": it["desc"], "size": 20, "bold": True, "color": self.card_ink, "space_after": 2})
        for extra in it["extra"] + it["children"]:
            lines.append({"text": extra, "size": 14, "color": self.card_muted})
        tb = self.text(s, cx + 0.3, cy + 0.2, cw - 0.6, ch - 0.4, lines, size, anchor=MSO_ANCHOR.MIDDLE)
        tb.rotation = rot

    def comp_compare(self, s, d, x, y, w, h):
        items = d["items"]
        n = len(items)
        if not n:
            return
        gap = 0.6 if n == 2 else 0.4
        cw = (w - gap * (n - 1)) / n
        hh = 0.95
        need = hh + 0.7 + max(self.lines_height(
            [{"text": "■ " + c, "size": 19, "space_after": 10} for c in it["children"] + it["extra"]], cw - 0.7)
            for it in items)
        ch = min(h, max(2.6, need))
        y += (h - ch) / 2
        h = ch
        toned = any(a in d["args"] for a in ("tone", "pros-cons", "proscons"))
        for k, it in enumerate(items):
            with self.tag(s, item=k, name=plain(it["title"])):
                self.compare_col(s, it, k, x + k * (cw + gap), y, cw, h, hh, toned)
        if n == 2 and not toned:
            self.star(s, x + cw + gap / 2 - 0.55, y + h / 2 - 0.55, 1.1, self.pal(3), text="VS", rot=10, size=16)

    def compare_col(self, s, it, k, cx, y, cw, h, hh, toned):
        hc = (self.tones["positive"] if k == 0 else self.tones["negative"]) if toned else self.color_for(k * 2)
        if self.tinted:
            self.tint_box(s, cx, y, cw, h, hc)
            self.text(s, cx + 0.3, y + 0.15, cw - 0.6, 0.5, it["title"], 22, bold=True, color=self.accent_ink(hc),
                      anchor=MSO_ANCHOR.MIDDLE, font="heading")
            self.shape(s, MSO_SHAPE.RECTANGLE, cx + 0.3, y + 0.68, min(cw - 0.6, self.measure(it["title"], 22) + 0.2),
                       0.05, fill=hc)
            if it["desc"]:
                self.text(s, cx + 0.3, y + 0.75, cw - 0.6, 0.3, it["desc"], 12, color=self.card_muted)
        else:
            self.block(s, cx, y, cw, h, mix(hc, self.ground, 0.85) if toned else self.white)
            self.shape(s, MSO_SHAPE.RECTANGLE, cx, y, cw, hh, fill=hc,
                       line=self.black if self.surface in ("brutal", "pixel") else None, lw=self.bw)
            head = [{"text": it["title"], "size": 24, "bold": True, "color": self.on(hc)}]
            if it["desc"]:
                head.append({"text": it["desc"], "size": 13, "color": self.on(hc)})
            self.text(s, cx, y, cw, hh, head, 24, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        lines = [{"text": "■ " + c, "size": 19, "color": self.card_ink, "space_after": 10}
                 for c in it["children"] + it["extra"]]
        if lines:
            self.text(s, cx + 0.35, y + hh + 0.35, cw - 0.7, h - hh - 0.5, lines, 19)

    def comp_chart(self, s, d, x, y, w, h):
        rows = d["rows"]
        if len(rows) < 2:
            self.warnings.append("chart block needs a Markdown table (first row = header)")
            return
        kind = (d["args"][0] if d["args"] else "column").lower()
        title = " ".join(d["args"][1:])
        types = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                 "stacked": XL_CHART_TYPE.COLUMN_STACKED, "line": XL_CHART_TYPE.LINE_MARKERS,
                 "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT}
        ct = types.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
        cd = CategoryChartData()
        cats = [r[0] for r in rows[1:]]
        cd.categories = cats

        def num(v):
            try:
                return float(re.sub(r"[,%\s]", "", v))
            except ValueError:
                return 0.0

        series_names = rows[0][1:]
        if kind in ("pie", "doughnut"):
            series_names = series_names[:1]
        for j, name in enumerate(series_names):
            cd.add_series(name, [num(r[j + 1]) if j + 1 < len(r) else 0 for r in rows[1:]])
        self.block(s, x, y, w, h, self.white)
        top = 0.25
        if title:
            self.pill(s, x + 0.25, y - 0.2, title, self.pal(3), size=12, rot=-2, upper=False, font="body")
            top = 0.4
        gf = s.shapes.add_chart(ct, E(x + 0.2), E(y + top), E(w - 0.4), E(h - top - 0.15), cd)
        ch = gf.chart
        ch.font.size = Pt(12)
        ch.font.name = self.f_body
        set_ea(ch.font._rPr, self.f_ea)
        multi = len(series_names) > 1 or kind in ("pie", "doughnut")
        ch.has_legend = multi
        if multi:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        plot = ch.plots[0]
        if kind in ("pie", "doughnut"):
            ser = plot.series[0]
            for i in range(len(cats)):
                pt = ser.points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(self.color_for(i))
                pt.format.line.color.rgb = rgb(self.black)
                pt.format.line.width = Pt(1.5)
        else:
            if kind in ("column", "bar", "stacked"):
                plot.gap_width = 60
                if kind == "stacked":
                    plot.overlap = 100
            for i, ser in enumerate(plot.series):
                c = self.color_for(i + 2 if len(plot.series) == 1 else i)
                if kind == "line":
                    ser.format.line.color.rgb = rgb(self.black if len(plot.series) == 1 else c)
                    ser.format.line.width = Pt(3)
                    ser.smooth = False
                    ser.marker.format.fill.solid()
                    ser.marker.format.fill.fore_color.rgb = rgb(c)
                    ser.marker.format.line.color.rgb = rgb(self.black)
                    ser.marker.size = 9
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = rgb(c)
                    ser.format.line.color.rgb = rgb(self.black)
                    ser.format.line.width = Pt(1.5)
            if len(cats) <= 8 and len(plot.series) <= 2:
                plot.has_data_labels = True
                plot.data_labels.font.size = Pt(11)
                plot.data_labels.font.bold = True
            va = ch.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = rgb("DDDDDD")
            va.format.line.fill.background()
            ch.category_axis.format.line.color.rgb = rgb(self.black)
            ch.category_axis.format.line.width = Pt(2)


def build_canvas(doc, theme, base_dir: Path):
    c = Canvas(theme, doc.meta, base_dir)
    prs = c.build(doc)
    return prs, c
