"""Structural QA for a built deck: text overflow, off-slide shapes, content outline.

`check()` returns a JSON-serialisable report so agents without image input can
still verify a deck; `format_report()` renders it for humans.
"""
from __future__ import annotations

import math
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from .layout import EMU_PER_PT
from .markdown import disp_width


def text_need(shp):
    """(needed, available) text height in pt for a shape's text frame, honouring its real margins."""
    tf = shp.text_frame
    w = (shp.width - tf.margin_left - tf.margin_right) / EMU_PER_PT
    h = (shp.height - tf.margin_top - tf.margin_bottom) / EMU_PER_PT
    need, last_sa = 0.0, 0.0
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        if not txt.strip():
            continue
        size = next((r.font.size.pt for r in p.runs if r.font.size), 18)
        last_sa = p.space_after.pt if p.space_after is not None else 0
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        pPr = p._p.pPr
        ind = int(pPr.get("marL", 0)) / EMU_PER_PT if pPr is not None else 0
        lines = sum(max(1, math.ceil(disp_width(seg) * size * 0.97 / max(20.0, w - ind))) for seg in txt.split("\n"))
        need += lines * size * 1.2 * ls + last_sa
    return need - last_sa, h


def check(path: str | Path) -> dict:
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    tol = Inches(0.15)
    slides, total = [], 0
    for n, slide in enumerate(prs.slides, 1):
        item = {"index": n, "layout": slide.slide_layout.name, "text": [], "pictures": 0, "tables": 0,
                "charts": 0, "notes": "", "issues": []}
        for shp in slide.shapes:
            if shp.left is not None and (shp.left < -tol or shp.top < -tol or shp.left + shp.width > W + tol
                                         or shp.top + shp.height > H + tol):
                item["issues"].append({"type": "off_slide", "shape": shp.name})
            if shp.has_text_frame:
                for p in shp.text_frame.paragraphs:
                    txt = "".join(r.text for r in p.runs)
                    if txt.strip():
                        item["text"].append({"level": p.level, "text": txt})
                need, avail = text_need(shp)
                if need > avail * 1.1 + 4:
                    item["issues"].append({"type": "text_overflow", "shape": shp.name,
                                           "needed_pt": round(need), "available_pt": round(avail),
                                           "excerpt": shp.text_frame.text[:30]})
            elif shp.shape_type == 13:
                item["pictures"] += 1
            elif getattr(shp, "has_chart", False) and shp.has_chart:
                item["charts"] += 1
            elif getattr(shp, "has_table", False) and shp.has_table:
                item["tables"] += 1
        if slide.has_notes_slide:
            item["notes"] = slide.notes_slide.notes_text_frame.text.strip()
        total += len(item["issues"])
        slides.append(item)
    return {"file": str(path), "slides": len(slides), "issues": total, "ok": total == 0, "pages": slides}


def format_report(report: dict, outline: bool = True) -> str:
    out = []
    for s in report["pages"]:
        head = next((t["text"] for t in s["text"]), "")
        out.append(f"── {s['index']:02d} [{s['layout']}] {head[:40]}")
        if outline:
            for t in s["text"][1:]:
                out.append(f"   {'  ' * t['level']}· {t['text']}")
            extras = [f"{k}={s[k]}" for k in ("pictures", "tables", "charts") if s[k]]
            if extras:
                out.append("   [" + ", ".join(extras) + "]")
            if s["notes"]:
                out.append(f"   [notes] {s['notes'][:60]}")
        for i in s["issues"]:
            if i["type"] == "text_overflow":
                out.append(f"   ⚠ text overflow: ~{i['needed_pt']}pt needed > {i['available_pt']}pt box — {i['excerpt']}")
            else:
                out.append(f"   ⚠ shape off slide: {i['shape']}")
    out.append(f"\n{report['slides']} slides, {report['issues']} issue(s)")
    return "\n".join(out)
