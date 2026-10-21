import json
from pathlib import Path

from deckframes.check import check
from deckframes.cli import main
from deckframes.frame_import import import_frame
from deckframes.markdown import parse_markdown

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "examples" / "demo.md"

FRAME = """---
name: test
colors:
  paper: "#F7F3EA"
  ink: "#151515"
  tomato: "#E4572E"
  teal: "#17BEBB"
  sun: "#FFC914"
borders: { primary: "4px solid black" }
shadows: { default: "8px 8px 0 black" }
typography:
  body:     { fontFamily: "Inter", cqw: 1.0 }
  headline:{ fontFamily: "Playfair Display", cqw: 5.0 }
---
"""


def test_hierarchy():
    doc = parse_markdown(DEMO.read_text(encoding="utf-8"))
    chapters = [s for s in doc.sections if s.title]
    assert len(chapters) == 4
    assert chapters[0].subtitle == "Introduction"
    assert chapters[0].subsections[0].slides[0].title == "消費習慣的轉變"
    kinds = {b.kind for s in chapters[0].subsections[0].slides for b in s.blocks}
    assert {"component", "source", "bullets"} <= kinds


def test_build_and_check(tmp_path):
    out = tmp_path / "demo.pptx"
    main(["build", str(DEMO), "-o", str(out)])
    report = check(out)
    assert report["slides"] >= 18
    assert report["ok"], json.dumps(report["pages"], ensure_ascii=False)[:500]


def test_canvas_master_is_widescreen(tmp_path):
    from pptx import Presentation
    out = tmp_path / "demo.pptx"
    main(["build", str(DEMO), "-o", str(out)])
    prs = Presentation(str(out))
    title = prs.slide_master.placeholders[0]
    # default python-pptx master is 4:3; its placeholders must be stretched to the 16:9 slide
    assert title.left + title.width > prs.slide_width * 0.9


def test_template_engine(tmp_path):
    out = tmp_path / "plain.pptx"
    tfile = tmp_path / "plain.json"
    tfile.write_text(json.dumps({"name": "plain", "fonts": {"heading": "Arial", "body": "Arial"}}), encoding="utf-8")
    main(["build", str(DEMO), "-o", str(out), "--theme", str(tfile)])
    assert check(out)["slides"] > 10


def test_frame_import(tmp_path):
    f = tmp_path / "FRAME.md"
    f.write_text(FRAME, encoding="utf-8")
    theme, _ = import_frame(f, "t")
    c = theme["colors"]
    assert c["ground"] == "F7F3EA" and c["text"] == "151515"
    assert set(c["palette"]) == {"E4572E", "17BEBB", "FFC914"}
    assert theme["decorations"] and theme["stroke"]["shadow"] == 4
    assert theme["fonts"]["display"] == "Georgia"
    tfile = tmp_path / "t.json"
    tfile.write_text(json.dumps(theme), encoding="utf-8")
    out = tmp_path / "t.pptx"
    main(["build", str(DEMO), "-o", str(out), "--theme", str(tfile)])
    assert check(out)["ok"]


def test_project_flow(tmp_path, monkeypatch):
    proj = tmp_path / "p"
    main(["init", str(proj), "--from", str(DEMO)])
    assert (proj / "deck.json").exists() and (proj / "AGENTS.md").exists()
    monkeypatch.chdir(proj)
    main(["build"])
    state = json.loads((proj / "deck.json").read_text(encoding="utf-8"))
    assert state["status"]["built"]
    assert (proj / state["output"]).exists()


def test_themes_new(tmp_path):
    from deckframes.themes import new_theme
    path = new_theme("my-lab", dest=tmp_path / "my-lab.json")
    theme = json.loads(path.read_text(encoding="utf-8"))
    assert theme["name"] == "my-lab" and theme["based_on"] == "blockframe" and theme["engine"] == "canvas"
    theme["colors"]["palette"] = ["1B998B", "ED217C", "FFFD82"]
    path.write_text(json.dumps(theme), encoding="utf-8")
    out = tmp_path / "t.pptx"
    main(["build", str(DEMO), "-o", str(out), "--theme", str(path)])
    assert check(out)["ok"]


def test_item_colon_and_workflow_alias(tmp_path):
    from deckframes.markdown import make_item
    assert make_item("空間老舊： | 座位不足")["title"] == "空間老舊"
    assert make_item("比例：3:1")["title"] == "比例：3:1"
    main(["init", str(tmp_path / "w"), "--workflow", "general"])
    state = json.loads((tmp_path / "w" / "deck.json").read_text(encoding="utf-8"))
    assert state["workflow"] == "deckframes-general"


def test_emoji_rejected(tmp_path):
    import pytest
    from deckframes.markdown import find_emoji
    assert [c for _, _, c in find_emoji("成果 \U0001F389 完成 \u2705 ok")] == ["\U0001F389", "\u2705"]
    assert find_emoji("✓ ✗ ○ × → ★ ■ – 3×") == []          # typographic symbols stay allowed
    assert [c for _, _, c in find_emoji("警告 \u26A0\uFE0F")] == ["⚠"]  # text symbol + VS16 = emoji
    deck = tmp_path / "e.md"
    deck.write_text("# 標題\n\n## 章 | Ch\n\n### 節\n\n- 重點 \U0001F680\n", encoding="utf-8")
    with pytest.raises(SystemExit) as e:
        main(["build", str(deck), "-o", str(tmp_path / "e.pptx")])
    assert e.value.code == 2 and not (tmp_path / "e.pptx").exists()


def test_check_flags_emoji_in_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Launch \U0001F680"
    out = tmp_path / "x.pptx"
    prs.save(out)
    issues = check(out)["pages"][0]["issues"]
    assert any(i["type"] == "emoji" for i in issues)


def test_icon_refs():
    from deckframes.icons import parse_ref
    assert parse_ref("users") == ("fa", None, "users")
    assert parse_ref("fa-solid fa-users") == ("fa", "solid", "users")
    assert parse_ref("regular:clock") == ("fa", "regular", "clock")
    assert parse_ref("brands:github") == ("fa", "brands", "github")
    assert parse_ref("assets/logo.svg") == ("file", "assets/logo.svg")
    assert parse_ref("?")[0] == "text" and parse_ref("AI")[0] == "text"


def test_svg_embedded_as_vector(tmp_path):
    import zipfile
    out = tmp_path / "d.pptx"
    main(["build", str(DEMO), "-o", str(out)])
    z = zipfile.ZipFile(out)
    assert any(n.endswith(".svg") for n in z.namelist())
    assert "image/svg+xml" in z.read("[Content_Types].xml").decode()


def _image_layout(tmp_path, size, body, title="#### 截圖 | 一句話說明這張圖", front=""):
    from PIL import Image
    from deckframes.engines.canvas import Canvas
    from deckframes.markdown import parse_markdown
    from deckframes.themes import resolve_theme
    Image.new("RGB", size, "white").save(tmp_path / "img.png")
    head = f"---\n{front}---\n\n" if front else ""
    doc = parse_markdown(f"{head}# T\n\n## 章 | Ch\n\n### 節\n\n{title}\n\n{body}\n")
    theme, _ = resolve_theme("blockframe")
    c = Canvas(theme, doc.meta, tmp_path)
    c.chapters = [s for s in doc.sections if s.title]
    sl = doc.sections[0].subsections[0].slides[0]
    from deckframes.engines.canvas import Page
    pg = Page("content", title=sl.title, subtitle=sl.subtitle, chapter=0, sub=0, blocks=sl.blocks)
    text, vis, callouts, sources = c.split_blocks(pg.blocks)
    return c.layout(pg, c.text_paras(text), vis, callouts, bool(sources))


def test_portrait_image_becomes_showcase(tmp_path):
    lay = _image_layout(tmp_path, (650, 1080), "![](img.png)")
    assert lay["mode"] == "showcase"
    assert lay["visual"][3] > 5.0                      # image column spans (almost) the full height


def test_text_and_image_split_follows_aspect(tmp_path):
    wide = _image_layout(tmp_path, (1754, 1241), "- 一點\n- 兩點\n\n![](img.png)")
    tall = _image_layout(tmp_path, (650, 1080), "- 一點\n- 兩點\n\n![](img.png)")
    assert wide["mode"] == tall["mode"] == "split"
    assert wide["visual"][2] > tall["visual"][2]       # landscape image gets the wider column
    left = _image_layout(tmp_path, (1754, 1241), '- 一點\n\n![](img.png "left")')
    assert left["visual"][0] < left["text"][0]         # "left" puts the image first


LONG = "- " + "A fairly long English sentence that would wrap many times in a narrow column. " * 3


def test_visual_below_puts_text_on_top(tmp_path):
    hint = _image_layout(tmp_path, (1754, 1241), '- one point\n\n![](img.png "below")')
    assert hint["mode"] == "below"
    assert hint["text"][1] < hint["visual"][1]         # text above the image
    assert hint["visual"][2] > 10                      # image spans the content width
    whole = _image_layout(tmp_path, (1754, 1241), "- one point\n\n![](img.png)", front="visual_below: true\n")
    assert whole["mode"] == "below"
    side = _image_layout(tmp_path, (1754, 1241), '- one point\n\n![](img.png "side")', front="visual_below: true\n")
    assert side["mode"] == "split"                     # "side" opts one slide out


def test_visual_below_auto_follows_wrapping(tmp_path):
    short = _image_layout(tmp_path, (1754, 1241), "- one point\n\n![](img.png)", front="visual_below: auto\n")
    long = _image_layout(tmp_path, (1754, 1241), f"{LONG}\n\n![](img.png)", front="visual_below: auto\n")
    assert short["mode"] == "split"                    # short text stays beside the image
    assert long["mode"] == "below"                     # text that keeps wrapping goes on top
    bled = _image_layout(tmp_path, (1754, 1241), f"{LONG}\n\n![](img.png)",
                         front="visual_below: auto\nvisual_bleed: 0.35\n")
    assert abs(bled["visual"][2] - long["visual"][2] - 0.7) < 1e-6   # bleed widens the image on both sides


COMPONENTS = ROOT / "examples" / "components.md"


def test_new_components_parse():
    doc = parse_markdown(COMPONENTS.read_text(encoding="utf-8"))
    from deckframes.lint import content_slides, slide_pattern

    used = {slide_pattern(sl.blocks) for sl in content_slides(doc)}
    assert {"lanes", "diagram", "mapping", "stack", "matrix", "funnel", "pyramid", "cycle", "progress"} <= used


def test_components_render_on_many_themes(tmp_path):
    for theme in ["academic", "aws", "blockframe"]:
        out = tmp_path / f"{theme}.pptx"
        try:
            main(["build", str(COMPONENTS), "--theme", theme, "-o", str(out)])
        except SystemExit as e:
            assert not e.code
        report = check(out)
        assert report["ok"], (theme, report["issues"][:3])


def test_gantt_columns_and_spans():
    from deckframes.markdown import gantt_columns, gantt_label, parse_component

    d = parse_component("gantt", "Jul'25-Sep'26 now=Sep", "group Study: Jul-Jan\n"
                        "- Draft | Plan | nov-May | window: Feb-May\n- Intern | | Jul'26-Aug'26\n"
                        "- Offer | | at: Sep'26")
    assert [gantt_label(c) for c in d["columns"]] == \
        "Jul Aug Sep Oct Nov Dec Jan Feb Mar Apr May Jun Jul Aug Sep".split()
    assert d["groups"] == [{"title": "Study", "span": (0, 6)}]
    assert d["items"][0]["span"] == (4, 10)
    assert d["items"][0]["window"] == (7, 10)
    assert d["items"][1]["span"] == (12, 13)
    assert d["items"][2]["at"] == (14, 14)
    assert len(gantt_columns(["Jan-Dec"])) == 12 and len(gantt_columns(["Oct-Mar"])) == 6
    for rng, expected in [("1-12", list(map(str, range(1, 13)))), ("2017-2026", list(map(str, range(2017, 2027))))]:
        d = parse_component("gantt", rng, f"- Work | Plan | {rng}")
        assert d["columns"] == expected
        assert d["items"][0]["span"] == (0, len(expected) - 1)
    d = parse_component("gantt", "cols=Q1,Q2,Q3,Q4", "- Work | Plan | Q2-Q3")
    assert d["columns"] == ["Q1", "Q2", "Q3", "Q4"] and d["items"][0]["span"] == (1, 2)


def test_steps_detail_panels_parse():
    from deckframes.markdown import parse_component

    d = parse_component("steps", "horizontal", "- Invitation | | Day + 0\n- Interviews\n"
                        "  - Prepare\n  > Technical\n    1. Project\n    2. Tradeoffs\n"
                        "  > Team\n    1. Collaboration")
    assert d["items"][0]["desc"] == "" and d["items"][0]["extra"] == ["Day + 0"]
    assert d["items"][1]["children"] == ["Prepare"]
    assert d["items"][1]["panels"] == [{"title": "Technical", "lines": ["Project", "Tradeoffs"]},
                                          {"title": "Team", "lines": ["Collaboration"]}]


def test_annotate_area_parse():
    from deckframes.markdown import parse_component

    d = parse_component("annotate", "img=assets/chart.png left build=focus",
                        "- Grade | Summary | area: 0,22,26,20\n- Project | Detail | area: 26,25,74,75")
    assert d["items"][0]["area"] == [0, 22, 26, 20]
    assert d["items"][1]["area"] == [26, 25, 74, 75]
    assert d["build"] == "focus" and d["args"] == ["img=assets/chart.png", "left"]


def test_annotate_focus_preserves_picture(tmp_path):
    from pptx import Presentation
    from pptx.oxml.ns import qn

    img = (ROOT / "examples/assets/chart.png").as_posix()
    src = tmp_path / "focus.md"
    src.write_text("---\noutline: false\nclosing: false\n---\n# Test\n## Part\n### Detail\n#### Image\n"
                   f"```annotate img={img} build=focus\n"
                   "- Summary | Overview | area: 0,0,100,25\n- Plot | Detail | area: 10,30,80,60\n```\n"
                   "+++ dim\n\nA final note.\n", encoding="utf-8")
    for theme in ["blockframe", "academic", "aws"]:
        out = tmp_path / f"focus-{theme}.pptx"
        main(["build", str(src), "--theme", theme, "-o", str(out)])
        slides = list(Presentation(str(out)).slides)[2:]
        assert len(slides) == 4
        for i, slide in enumerate(slides[:3]):
            names = {sh.name for sh in slide.shapes}
            regions = {name for name in names if name.startswith("annotate-region-")}
            assert regions == ({"annotate-region-1", "annotate-region-2"} if i == 0 else {f"annotate-region-{i}"})
            frames = {name for name in names if name.startswith("annotate-card-")}
            assert frames == (set() if i == 0 else {f"annotate-card-{i}"})
            picture = next(sh for sh in slide.shapes if sh._element.tag == qn("p:pic"))
            assert not list(picture._element.iter(qn("a:alphaModFix")))
        picture = next(sh for sh in slides[-1].shapes if sh._element.tag == qn("p:pic"))
        assert list(picture._element.iter(qn("a:alphaModFix")))
        assert check(out)["ok"]


def test_presentation_patterns_reveal_on_builtin_themes(tmp_path):
    from pptx import Presentation
    from pptx.oxml.ns import qn

    assets = (ROOT / "examples/assets").as_posix()
    src = tmp_path / "patterns.md"
    src.write_text(f"""---
outline: false
closing: Q&A
closing_image: {assets}/photo.png
---
# Patterns
## Part
### Details
#### Roadmap
```gantt 1-12 build=dim
group Plan: 1-6
group Deliver: 7-12
- Research | First draft | 1-5
- Launch | Release | at: 12 | note: Ready
```
#### Annotation
```annotate img={assets}/chart.png left build
- Overview | Read the summary | area: 0,0,100,20
  - Compare the series
- Evidence | Check the values | area: 0,20,50,80
- Detail | Inspect the scale | area: 50,20,50,80
```
#### Process
```steps horizontal build=focus
- Invitation | | Day + 0
- Interviews
  - Prepare a demo
  > Technical
    1. Project
    2. Tradeoffs
  > Team
    1. Collaboration
- Review
- Decision
```
#### Follow-up
```steps vertical
- Interview
  > Discussion
    1. Share a project
- Decision
```
#### Day
```timeline axis build
- 08:30 | Arrive | By bus | img: {assets}/photo.png
- 12:00 | Lunch | Shared break | img: {assets}/photo.png
```
#### Partners
```stats build=focus
- Invitations | img: {assets}/brand-a.svg, {assets}/brand-a.svg
- Pilot | img: {assets}/brand-a.svg
```
> [!HINT] Keep a working example ready.
""", encoding="utf-8")
    for theme in ["blockframe", "academic", "aws"]:
        out = tmp_path / f"patterns-{theme}.pptx"
        main(["build", str(src), "--theme", theme, "-o", str(out)])
        report = check(out)
        assert report["ok"], (theme, report["issues"][:3])
        slides = list(Presentation(str(out)).slides)
        annotations = [s for s in slides if any(sh.name.startswith("annotate-region-") for sh in s.shapes)]
        assert [sum(sh.name.startswith("annotate-region-") for sh in s.shapes) for s in annotations] == [1, 2, 3]
        for s in slides:
            text = "\n".join(sh.text for sh in s.shapes if sh.has_text_frame)
            assert text.count("Day + 0") <= 1
        closing = slides[-1]
        pic = next(sh for sh in closing.shapes if sh._element.tag == qn("p:pic"))
        prs = Presentation(str(out))
        assert (pic.left, pic.top, pic.width, pic.height) == (0, 0, prs.slide_width, prs.slide_height)
        assert pic.crop_top > 0 and pic.crop_top == pic.crop_bottom


def test_builtin_themes_have_distinct_styles():
    from deckframes.themes import list_themes

    themes = [json.loads(Path(t["path"]).read_text(encoding="utf-8")) for t in list_themes()
              if t["engine"] == "canvas" and t["source"] == "builtin"]
    assert {t["name"] for t in themes} == {"academic", "aws", "blockframe"}
    sigs = {(t["style"]["surface"], t["style"].get("cover"), t["style"].get("nav")) for t in themes if "style" in t}
    assert len(sigs) == len(themes)


def test_monotony_warning():
    from deckframes.lint import variety

    slides = "\n".join(f"### Slide {i}\n\n- a\n- b\n" for i in range(6))
    doc = parse_markdown(f"# Deck\n\n## Part\n\n{slides}")
    warns = variety(doc)
    assert any("consecutive" in w for w in warns)
    assert any("6/6" in w for w in warns)
