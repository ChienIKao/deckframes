"""Build steps (`+++`), item focus / reveal, overlays and the tinted card style."""
from pathlib import Path

from pptx import Presentation

from deckframes.check import check
from deckframes.cli import main
from deckframes.markdown import parse_markdown

ROOT = Path(__file__).resolve().parents[1]
BUILDS = ROOT / "examples" / "builds.md"
HEAD = "---\ntheme: academic\noutline: false\nclosing: false\n---\n\n# T\n\n## 章 | Ch\n\n### 節\n\n"


def _deck(tmp_path, body, name="d"):
    src = tmp_path / f"{name}.md"
    src.write_text(HEAD + body, encoding="utf-8")
    out = tmp_path / f"{name}.pptx"
    main(["build", str(src), "-o", str(out)])
    return Presentation(str(out)), out


def _content(prs):
    """Content slides only (skip cover + chapter divider)."""
    return list(prs.slides)[2:]


def _texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def _box(slide, text):
    sh = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.startswith(text))
    return sh.left, sh.top, sh.width, sh.height


def _run_color(slide, text):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.startswith(text):
                        return str(r.font.color.rgb)
    raise AssertionError(text)


def test_step_markers_parse():
    doc = parse_markdown(HEAD + "#### S\n\n- a\n+++ dim\n<!-- second -->\n- b\n\n+++ overlay focus=2,3\n\n"
                         "```cards build=focus cols=2\n- x\n- y\n```\n")
    blocks = doc.sections[0].subsections[0].slides[0].blocks
    steps = [b.data for b in blocks if b.kind == "step"]
    assert steps[0]["dim"] and not steps[0]["overlay"] and steps[0]["notes"] == ["second"]
    assert steps[1]["overlay"] and steps[1]["focus"] == "2,3"
    comp = next(b.data for b in blocks if b.kind == "component")
    assert comp["build"] == "focus" and comp["args"] == ["cols=2"]   # build/focus never leak into args


def test_bullets_build_keeps_positions(tmp_path):
    prs, out = _deck(tmp_path, "#### 逐點 | sub\n\n- 第一點\n+++\n- 第二點\n+++\n- 第三點\n")
    slides = _content(prs)
    assert len(slides) == 3
    assert "第三點" not in "".join(_texts(slides[1])) and "第三點" in "".join(_texts(slides[2]))
    assert _box(slides[0], "第一點") == _box(slides[2], "第一點")    # laid out once for all steps
    assert check(out)["ok"]


def test_component_build_reveals_items(tmp_path):
    prs, _ = _deck(tmp_path, "#### 卡片\n\n```steps build\n- 甲 | 一\n- 乙 | 二\n- 丙 | 三\n```\n")
    slides = _content(prs)
    assert len(slides) == 3
    assert [sum(t.startswith(("甲", "乙", "丙")) for t in _texts(s)) for s in slides] == [1, 2, 3]
    assert _box(slides[0], "甲") == _box(slides[2], "甲")


def test_diagram_build_reveals_every_node(tmp_path):
    prs, _ = _deck(tmp_path, "#### 架構\n\n```diagram build focus=甲\n- 甲 -> 乙\n- 乙 -> 丙\n- 丁 -> 丙\n```\n\n"
                             "+++ focus=丙\n\n> [!NOTE] 丙\n")
    slides = _content(prs)
    assert len(slides) == 5                                            # 4 nodes (not 3 edge lines) + focus step
    assert any(t.startswith("丁") for t in _texts(slides[3]))
    assert _run_color(slides[4], "丙") != _run_color(slides[4], "甲")   # the step's focus overrides focus=甲


def test_focus_greys_out_the_other_items(tmp_path):
    prs, _ = _deck(tmp_path, "#### 對照\n\n```mapping focus=2\n- 問題一 -> 解法一\n- 問題二 -> 解法二\n```\n")
    s = _content(prs)[0]
    assert _run_color(s, "問題一") != _run_color(s, "問題二")         # item 1 faded, item 2 keeps its ink


def test_dim_step_and_build_focus(tmp_path):
    prs, _ = _deck(tmp_path, "#### 流程\n\n```flow\n- 收到\n- 比對\n```\n\n+++ dim\n\n```flow\n- 排序\n```\n")
    first, second = _content(prs)
    assert _run_color(first, "收到") != _run_color(second, "收到")
    assert _run_color(second, "排序") == _run_color(first, "收到")
    prs, _ = _deck(tmp_path, "#### 走查\n\n```cards build=focus\n- 甲\n- 乙\n```\n", name="f")
    assert len(_content(prs)) == 3                                     # overview + one slide per item


def test_overlay_veils_the_slide_below(tmp_path):
    body = ("#### 表格\n\n| 模型 | 分數 |\n|---|---|\n| A | 1 |\n| B | 2 |\n\n"
            "+++ overlay focus=2\n\n> [!RESULT] B 最好\n")
    prs, out = _deck(tmp_path, body)
    base, over = _content(prs)
    assert not any(sh.name == "veil" for sh in base.shapes)
    names = [sh.name for sh in over.shapes]
    assert "veil" in names
    title_at = next(i for i, sh in enumerate(over.shapes) if sh.has_text_frame and sh.text_frame.text.startswith("節 – 表格"))
    assert title_at > names.index("veil")                              # the title is lifted above the veil
    assert check(out)["ok"]


def test_builds_example_and_template_fallback(tmp_path):
    out = tmp_path / "b.pptx"
    main(["build", str(BUILDS), "-o", str(out)])
    report = check(out)
    assert report["ok"] and report["slides"] >= 25
    import json
    tfile = tmp_path / "plain.json"
    tfile.write_text(json.dumps({"name": "plain", "fonts": {"heading": "Arial", "body": "Arial"}}), encoding="utf-8")
    flat = tmp_path / "t.pptx"
    main(["build", str(BUILDS), "-o", str(flat), "--theme", str(tfile)])   # templates show the finished slides
    assert check(flat)["slides"] < report["slides"]


def test_tinted_cards_use_accent_outlines(tmp_path):
    prs, _ = _deck(tmp_path, "#### 卡片\n\n```cards\n- 甲 | 一\n- 乙 | 二\n```\n")
    s = _content(prs)[0]
    lines = {str(sh.line.color.rgb) for sh in s.shapes
             if sh.shape_type == 1 and sh.line.fill.type == 1 and sh.top > 914400 * 1.5}
    assert "E9A23B" in lines                                           # academic: gold card outlines
