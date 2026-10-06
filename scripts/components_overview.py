"""Render the example decks and tile one labelled slide per component into docs/components.jpg.

    python scripts/components_overview.py [--out docs/components.jpg] [--backend auto|powerpoint|libreoffice]

Slides are picked by a text snippet that appears on exactly one slide of the deck (an int picks by
position, -1 = last), so the overview survives slides being added to the examples.
"""
from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation

from deckframes.gallery import _font
from deckframes.preview import render

ROOT = Path(__file__).resolve().parent.parent
DECKS = {"c": "examples/components.md", "d": "examples/demo.md", "s": "examples/sample.md"}

GROUPS = [
    ("頁面版型", [
        ("c", 1, "封面", "標題、副標、作者、日期"),
        ("c", 2, "大綱", "所有章節串成一條線"),
        ("c", 3, "章節頁", "章節色面板＋子章節列表"),
        ("c", -1, "結尾", "Q&A / 謝謝聆聽"),
    ]),
    ("列舉與對照", [
        ("d", "門市三大痛點", "cards", "痛點、重點，各配一個圖示"),
        ("d", "消費習慣的轉變", "cards stack", "並列案例，可附標籤與圖片"),
        ("c", "問題與解法對照", "mapping", "問題 → 解法逐一對應"),
        ("d", "點餐工具比較", "表格 + [!NOTE]", "方法 × 特性矩陣（○ / ×）"),
    ]),
    ("流程與時間", [
        ("d", "資料流程", "flow", "系統流程、管線"),
        ("d", "經營模式的轉變", "flow vertical", "轉變：舊 → 新"),
        ("c", "前處理四步驟", "steps + [!TIP]", "方法步驟（橫式）"),
        ("c", "前處理檢查點", "steps vertical", "方法步驟（直式）"),
        ("d", "四個階段", "timeline", "階段時程"),
        ("c", "研究循環", "cycle", "反覆循環，中心放標籤"),
        ("c", "兩條研究路線", "lanes", "依年份排的多條研究脈絡"),
    ]),
    ("架構與層級", [
        ("c", "任務定義", "diagram", "有分支的架構圖、資料流"),
        ("c", "整體架構", "diagram focus= + group", "凸顯一部分，其餘淡化"),
        ("c", "問題一：看不出區域差異", "diagram focus= + [!PROBLEM]", "連續幾頁逐一講解"),
        ("c", "網路層級", "stack", "模型層、技術堆疊"),
        ("c", "研究層次", "pyramid", "階層，越上越窄"),
        ("c", "資料處理漏斗", "funnel", "逐步篩選"),
        ("c", "方法比較矩陣", "matrix", "2×2 定位"),
    ]),
    ("比較與數字", [
        ("d", "為何選擇漸進導入", "compare", "方案並排，兩欄有 VS"),
        ("c", "對我們研究的意義", "compare tone", "優缺點（綠 / 紅欄）"),
        ("d", "排隊時間縮短", "stats + [!IMPORTANT]", "大字數字重點"),
        ("c", "各時間尺度誤差", "chart + [!RESULT]", "原生可編輯圖表"),
        ("c", "實作進度", "progress", "完成度"),
    ]),
    ("文字與圖片", [
        ("d", "提案範圍與限制", "條列", "沒有視覺元件時的預設"),
        ("d", "的點餐規律", "==重點== 段落", "宣言頁"),
        ("d", "讓每一杯咖啡", "> 引言", "引言卡片"),
        ("s", "訪談 42 位", "圖片", "文字左、圖右（依比例排版）"),
        ("d", "報表匯出排程", "程式碼", "程式碼區塊"),
    ]),
]

COLS, TW, TH = 4, 640, 360
GAP, PAD, LABEL_H, HEAD_H, TITLE_H = 28, 48, 78, 84, 90
BG, INK, MUTED, EDGE = (246, 244, 238), (20, 20, 20), (110, 110, 110), (190, 190, 190)


def slide_texts(pptx: Path) -> list[str]:
    return ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
            for s in Presentation(str(pptx)).slides]


def locate(texts: list[str], key, deck: str) -> int:
    """1-based slide number for a position or a snippet that occurs on exactly one slide."""
    if isinstance(key, int):
        return key if key > 0 else len(texts) + 1 + key
    hits = [i + 1 for i, t in enumerate(texts) if key in t]
    if len(hits) != 1:
        raise SystemExit(f"{deck}: {key!r} matches slides {hits or 'none'}, expected exactly one")
    return hits[0]


def build_decks(tmp: Path, backend: str) -> dict[str, tuple[Path, list[str]]]:
    rendered = {}
    for tag, md in DECKS.items():
        pptx = tmp / f"{tag}.pptx"
        subprocess.run([sys.executable, "-m", "deckframes", "build", str(ROOT / md), "-o", str(pptx)],
                       check=True, stdout=subprocess.DEVNULL)
        out = Path(render(pptx, tmp / f"{tag}_preview", backend=backend)["dir"])
        rendered[tag] = (out, slide_texts(pptx))
    return rendered


def compose(rendered, out: Path) -> tuple[int, int]:
    f_title, f_head = _font(52, bold=True), _font(34, bold=True)
    f_name, f_desc = _font(24, bold=True), _font(20)
    width = PAD * 2 + COLS * TW + (COLS - 1) * GAP
    rows = sum(-(-len(items) // COLS) for _, items in GROUPS)
    height = PAD + TITLE_H + len(GROUPS) * HEAD_H + rows * (TH + LABEL_H + GAP) + PAD
    img = Image.new("RGB", (width, height), BG)
    d = ImageDraw.Draw(img)
    d.text((PAD, PAD), "deckframes 元件總覽", font=f_title, fill=INK)
    y = PAD + TITLE_H
    for head, items in GROUPS:
        d.text((PAD, y + 22), head, font=f_head, fill=INK)
        d.line((PAD, y + HEAD_H - 10, width - PAD, y + HEAD_H - 10), fill=INK, width=3)
        y += HEAD_H
        for i, (tag, key, name, desc) in enumerate(items):
            r, c = divmod(i, COLS)
            x, ty = PAD + c * (TW + GAP), y + r * (TH + LABEL_H + GAP)
            folder, texts = rendered[tag]
            n = locate(texts, key, DECKS[tag])
            with Image.open(folder / f"slide-{n:03d}.png") as im:
                img.paste(im.convert("RGB").resize((TW, TH), Image.LANCZOS), (x, ty))
            d.rectangle((x - 1, ty - 1, x + TW, ty + TH), outline=EDGE, width=1)
            d.text((x, ty + TH + 12), name, font=f_name, fill=INK)
            d.text((x, ty + TH + 44), desc, font=f_desc, fill=MUTED)
        y += -(-len(items) // COLS) * (TH + LABEL_H + GAP)
    out.parent.mkdir(parents=True, exist_ok=True)
    img.save(out, quality=84, optimize=True, progressive=True)
    return img.size


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", default=str(ROOT / "docs" / "components.jpg"))
    ap.add_argument("--backend", default="auto", choices=["auto", "powerpoint", "libreoffice"])
    a = ap.parse_args()
    with tempfile.TemporaryDirectory() as tmp:
        size = compose(build_decks(Path(tmp), a.backend), Path(a.out))
    print(f"✔ {a.out}  {size[0]}×{size[1]}")


if __name__ == "__main__":
    main()
